from pathlib import Path

from docx import Document


for path in [
    Path("artifacts/Smart_Water_Monitoring_Consumption_System_Chapter_4.docx"),
    Path("artifacts/Smart_Water_Monitoring_Consumption_System_Chapters_4_and_5.docx"),
]:
    document = Document(path)
    print(f"FILE {path} paragraphs={len(document.paragraphs)} tables={len(document.tables)} shapes={len(document.inline_shapes)}")
    for index, paragraph in enumerate(document.paragraphs, 1):
        text = " ".join(paragraph.text.split())
        if text:
            print(f"{index}|{paragraph.style.name}|{text[:280]}")
