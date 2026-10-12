from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\laragon\capstone\SmartWater_Update\Smart-Water-Consumption-Monitoring-Revised.docx")
OUTPUT = Path(r"C:\laragon\capstone\SmartWater_Update\Smart-Water-Consumption-Monitoring-Updated-Chapters-I-III.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = tc_pr.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        tc_pr.append(shade)
    shade.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    set_keep_with_next(p)
    return p


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.35)
    p.paragraph_format.space_after = Pt(6)
    p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(3)
    p.add_run("• " + text)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(3)
    p.add_run("• " + text)
    return p


def add_chapter_label(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    return p


def add_simple_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, (label, width) in enumerate(zip(headers, widths)):
        cell = header.cells[i]
        cell.width = Inches(width)
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)
    for index, row_data in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for i, (value, width) in enumerate(zip(row_data, widths)):
            cell = cells[i]
            cell.width = Inches(width)
            set_cell_margins(cell)
            if index % 2:
                set_cell_shading(cell, "F2F6F9")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            r.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


doc = Document(SOURCE)

# Retain the source document's styles, page geometry, and header/footer, but replace
# its body with the updated Chapters I–III content.
body = doc._element.body
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)

for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles["Normal"].font.size = Pt(12)
doc.styles["Normal"].paragraph_format.line_spacing = 1.5
doc.styles["Normal"].paragraph_format.space_after = Pt(6)
for level, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
    doc.styles[level].font.size = Pt(size)
    doc.styles[level].font.bold = True

# Cover page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(120)
p.paragraph_format.space_after = Pt(24)
r = p.add_run("SMART WATER CONSUMPTION\nMONITORING SYSTEM")
r.bold = True
r.font.name = "Times New Roman"
r.font.size = Pt(18)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(54)
r = p.add_run("CHAPTERS I TO III")
r.font.size = Pt(14)
r.bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(68)
r = p.add_run("Updated system documentation\nCalape, Bohol")
r.font.size = Pt(12)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Aligned with the implemented Laravel and ESP32 project\nOctober 2026")
r.font.size = Pt(12)
doc.add_page_break()

# Chapter I
add_chapter_label(doc, "CHAPTER I")
add_heading(doc, "INTRODUCTION", 1)
add_heading(doc, "1.1 Background of the Project", 2)
add_body(doc, "Households commonly learn their water consumption only after an official billing statement is issued. That delay makes it difficult to relate day-to-day water use to cost, investigate unusual consumption promptly, or distinguish a possible plumbing problem from normal use. Smart metering addresses this visibility gap by collecting consumption data at a shorter interval and presenting it in a form that households and water-service personnel can act upon. Research on smart water systems similarly identifies real-time consumption visibility and timely loss information as practical benefits of connected metering.")
add_body(doc, "The Smart Water Consumption Monitoring System is a household-scale Internet of Things project for Calape, Bohol. A YF-S201 flow sensor produces pulse signals as water passes through the monitored pipeline. An M5Stack CoreS3, built on the ESP32-S3 platform, samples those pulses, estimates interval flow and volume, displays the current state locally, and transmits interval readings through Wi-Fi to a Laravel web application. The application preserves readings, derives the server-authoritative monthly total, presents household and municipal dashboards, and supports administrative monitoring, estimated billing, notifications, reports, and audit review.")
add_body(doc, "The updated implementation does more than display a device total. A meter is first registered in a municipal device registry, commissioned with a one-time activation code, and bound to a household. Each uploaded reading has a stable identifier so a locally retained interval can be retried after a connection interruption without being counted twice. The server, rather than the device payload, determines the accepted household assignment, consumption period total, applicable billing settings, and possible-leak status. This design reduces the risk that a temporary connectivity problem or repeated transmission will distort the record.")
add_heading(doc, "1.2 Statement of the Problem", 2)
add_body(doc, "The project addresses the following operational problems in household water monitoring and municipal support:")
for item in [
    "Households have limited visibility of current consumption before the next official billing cycle.",
    "Abnormal or continuous flow may remain unnoticed until it causes higher consumption or visible damage.",
    "Water-service personnel need a controlled way to register meters, associate them with households, review consumption, manage billing settings, and retain an audit trail.",
    "A network interruption can cause missing or duplicate telemetry unless the device and server use a safe retry and idempotency design.",
    "Estimated monitoring information must not be confused with the municipality's official metering and billing process until calibration and local procedures are formally validated."
]:
    add_numbered(doc, item)
add_heading(doc, "1.3 Purpose and Description of the Project", 2)
add_body(doc, "The purpose of the project is to provide a web-connected household water-consumption monitoring system that gives households timely visibility of measured use while supplying authorized municipal personnel with role-scoped operational records. The system combines a locally visible meter panel with a web application. It is intended to support awareness, monitoring, and follow-up; it does not replace an approved utility meter, an official bill, or a physical inspection.")
add_body(doc, "At the device level, the CoreS3 samples the YF-S201 pulse signal on GPIO2, calculates interval usage from a configurable calibration factor, and uploads a reading approximately every three seconds when connectivity is available. Pending readings are preserved in device storage. At the application level, Laravel validates device credentials and the device-to-household assignment, stores accepted readings, calculates month-to-date consumption, evaluates configured leak conditions, and prepares billing-related information using the applicable rate, fixed charge, and subsidy or deduction settings.")
add_heading(doc, "1.4 Objectives", 2)
add_heading(doc, "General Objective", 3)
add_body(doc, "To develop an IoT-based Smart Water Consumption Monitoring System that measures household water flow, transmits reliable interval telemetry to a web application, and provides role-scoped consumption, alert, and estimated-billing information.")
add_heading(doc, "Specific Objectives", 3)
for item in [
    "Measure water-flow pulses from a YF-S201 sensor and convert calibrated pulse counts into interval liters and flow rate using an M5Stack CoreS3.",
    "Display local device status, including measured flow, current monitoring state, Wi-Fi status, and the server-authoritative period total when available.",
    "Commission each physical meter through a municipal registry, one-time activation code, hardware identity, and scoped device credential.",
    "Transmit readings through Wi-Fi using stable reading identifiers and retain unsent intervals for safe retry after interruption.",
    "Provide household access to live and historical consumption, estimated and finalized billing information, notifications, usage goals, service requests, and downloadable billing documents.",
    "Provide administrators with household, device, billing, leak, notification, reporting, payment-ledger, service-request, and audit-log functions.",
    "Evaluate possible leaks from server-configured flow-rate and duration conditions, and create notifications and follow-up records without claiming that the device can physically shut off water.",
    "Verify the implemented software functions through automated tests while identifying physical calibration and field validation as separate requirements."
]:
    add_numbered(doc, item)
add_heading(doc, "1.5 Scope and Limitations", 2)
add_heading(doc, "Scope", 3)
add_simple_table(doc, ["Area", "Implemented Scope"], [
    ("Physical meter", "A YF-S201 pulse-output flow sensor connected to an M5Stack CoreS3. The firmware samples flow once per second and attempts telemetry upload every three seconds."),
    ("Household workspace", "Live state, consumption history and trends, estimated and finalized statements, notifications, goals, account settings, low-bandwidth mode, and municipal service requests."),
    ("Municipal workspace", "Household approval and maintenance, device registry and lifecycle actions, configurable billing settings, leak monitoring, reports and exports, offline-payment records, action center, and audit logs."),
    ("Data integrity", "One-time device activation, scoped credentials, stable reading identifiers, server-side deduplication, device retry persistence, and server-authoritative consumption and alert decisions."),
    ("Billing support", "Estimated and finalized statement calculations using effective rates, fixed monthly charges, subsidies or deductions, opening balance, and recorded offline payments.")
], [1.6, 4.9])
add_heading(doc, "Limitations", 3)
for item in [
    "The initial firmware calibration value is 450 pulses per liter. It is a starting configuration, not a verified field accuracy claim. A measured-volume calibration test is still required for each installed sensor and plumbing arrangement.",
    "The system requires electrical power and usable Wi-Fi for timely upload. Pending intervals can be retried, but a prolonged device, power, network, or server outage can delay the dashboard view.",
    "The system is designed for household monitoring and municipal support. It has not been validated for industrial, commercial, or utility-wide production deployment.",
    "Possible-leak status is a server-side alert based on configured flow conditions. It is not a physical diagnosis and does not actuate a valve or automatically stop water.",
    "Application estimates and system-generated statements are decision-support records. They must be reconciled with approved municipal rules and official billing procedures before being treated as official bills.",
    "Automated software tests establish coverage of implemented behaviors, but they do not establish sensor accuracy, installation quality, long-term reliability, backup restoration readiness, or unassisted usability by real household participants."
]:
    add_bullet(doc, item)
add_heading(doc, "1.6 Significance of the Project", 2)
add_body(doc, "For households, the system offers a more timely view of consumption and a way to investigate an unusual flow or changing monthly pattern before the next official bill. For municipal personnel, it organizes device identity, household assignment, consumption records, billing settings, reports, and follow-up activity in one access-controlled workspace. For the project team, it provides a practical integration of embedded sensing, networked telemetry, web application development, data integrity controls, and testable service logic.")
add_body(doc, "The project is also relevant to future researchers because it documents a constrained household prototype rather than claiming an already proven utility replacement. The separation between application verification and physical validation makes clear which evidence is available now and what must still be gathered in a pilot installation.")
add_heading(doc, "1.7 Definition of Terms", 2)
for term, definition in [
    ("Commissioning", "The controlled process that activates a registered physical meter and exchanges a one-time code for a scoped device credential."),
    ("Idempotent reading", "A telemetry record with a stable identifier that the server accepts only once, even when the device resends it after an interruption."),
    ("Server-authoritative total", "The monthly consumption total calculated from accepted server records rather than a cumulative value supplied by the device."),
    ("Possible leak", "A system alert produced when configured continuous-flow conditions are met; it requires inspection and is not proof of a plumbing fault."),
    ("Estimated bill", "A calculated consumption charge based on the configured rate, fixed charge, and subsidy or deduction. It is not automatically an official municipal bill.")
]:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(term + ". ")
    r.bold = True
    p.add_run(definition)

doc.add_page_break()

# Chapter II
add_chapter_label(doc, "CHAPTER II")
add_heading(doc, "REVIEW OF RELATED LITERATURE AND SYSTEMS", 1)
add_heading(doc, "2.1 Introduction", 2)
add_body(doc, "This chapter reviews work relevant to household-scale water monitoring, connected flow sensing, consumption visualization, and loss detection. The review is used to position the present project, not to claim equivalence with the cited systems. The current project differs in its use of a M5Stack CoreS3 and YF-S201 sensor, a Laravel municipal device lifecycle, server-side idempotent telemetry handling, and role-scoped operational records.")
add_heading(doc, "2.2 Real-Time Water Consumption Monitoring", 2)
add_body(doc, "Andrić et al. (2022) describe an IoT approach to smart water use in which connected meters support more transparent consumption information and analysis. Their work is relevant because it treats metering as a data path from physical measurement to understandable consumption information, rather than as a display-only device. The present project applies the same broad principle at household scale, but its implementation uses Wi-Fi telemetry and a web-based municipal workflow rather than the communication environment of that study.")
add_body(doc, "A visual-analytics water-monitoring system reported by Tasong and Abao (2019) used Hall-effect flow sensing, cloud storage, browser access, and calibration trials. It is particularly relevant to the present project because it demonstrates that pulse-derived volume requires calibration rather than assumption. In response, the present documentation identifies the firmware factor as an initial configuration and treats measured-volume calibration as a remaining validation activity, not as a completed accuracy result.")
add_heading(doc, "2.3 Smart Metering, Loss Awareness, and Missing Data", 2)
add_body(doc, "High-resolution smart-meter data can support water-loss assessment, but incomplete or unavailable readings can affect operational interpretation. The work of Bragalli, Neri, and Toth (2019) is relevant because it explicitly considers the effect of missing readings in a real network. The present system therefore retains unsent intervals locally, uses a stable reading identifier on retry, and prevents the server from counting the same device reading twice. These controls improve resilience of the application record, although they do not eliminate all risks from a prolonged outage or failed physical device.")
add_body(doc, "The project also uses a possible-leak alert rather than a conclusive leak diagnosis. A configurable flow-rate threshold and continuous-duration condition are evaluated by the server; when the conditions are satisfied, a household notification and leak record can be created. This approach supports early follow-up while maintaining the appropriate boundary that a sensor-based anomaly signal still requires on-site inspection.")
add_heading(doc, "2.4 Secure IoT Device Lifecycle", 2)
add_body(doc, "A household meter becomes an operational system only when its data can be attributed to the correct physical device and household. The present project implements a municipal device registry with asset metadata, registration, assignment, activation, suspension, credential rotation, and decommissioning. During commissioning, a one-time activation code and hardware identifier are exchanged for a scoped credential. This design is consistent with token-based API concepts documented by Laravel Sanctum, while the specific activation and lifecycle rules are implemented in the project itself.")
add_body(doc, "This lifecycle distinction is important because a generic shared upload key would not provide the same accountability. The project's approach links accepted telemetry to a registered device, preserves a firmware version and last-seen information, and allows a credential to be revoked if the meter is replaced, suspended, or decommissioned.")
doc.add_page_break()
add_heading(doc, "2.5 Synthesis of the Review", 2)
add_simple_table(doc, ["Related Finding", "Implication for This Project", "Current Treatment"], [
    ("Real-time consumption information can improve visibility of household use.", "Present data in a timely, understandable household view.", "Live state, history, trends, statements, notifications, and low-bandwidth access."),
    ("Pulse-based sensing requires calibration for credible volume values.", "Avoid making accuracy claims before measured-volume testing.", "Initial pulse factor is configurable; physical calibration remains a documented limitation."),
    ("Smart-meter readings can be incomplete or repeated during communication failures.", "Design for retry without double counting.", "Pending device storage, stable reading IDs, and server-side idempotency."),
    ("Anomaly data supports early investigation but does not replace inspection.", "Use alerts as operational prompts, not final diagnoses.", "Server-side possible-leak evaluation, notifications, and follow-up lifecycle."),
    ("Connected devices require controlled identity and access.", "Bind telemetry to a registered and revocable device identity.", "One-time activation, hardware binding, scoped credential, and lifecycle actions.")
], [1.75, 2.3, 2.45])
add_heading(doc, "2.6 Research Gap Addressed", 2)
add_body(doc, "Existing work demonstrates the feasibility of connected water sensing, visual consumption information, and anomaly-oriented monitoring. The present project addresses a practical gap between a prototype meter and a local operational workflow. It adds an explicit municipal registry, household assignment, retry-safe telemetry, server-authoritative consumption, configurable billing support, managed leak follow-up, reports, and audit trails. The contribution is therefore not a claim of a new sensing principle; it is an integrated and testable household-monitoring workflow tailored to the implemented system.")

doc.add_page_break()

# Chapter III
add_chapter_label(doc, "CHAPTER III")
add_heading(doc, "TECHNICAL BACKGROUND", 1)
add_heading(doc, "3.1 System Architecture", 2)
add_body(doc, "The Smart Water Consumption Monitoring System has four connected layers: measurement, device processing, application services, and user workspaces. The measurement layer is the YF-S201 sensor installed in a household water path. The device layer is the M5Stack CoreS3, which receives the pulse signal through GPIO2, converts pulse counts using its calibration setting, maintains a local pending-reading state, and presents essential status information on its screen. The application layer is a Laravel service with a relational database, APIs, business services, and audit logging. The user layer consists of household and administrator browser workspaces with access restricted by role and record ownership.")
add_simple_table(doc, ["Layer", "Primary Responsibility", "Examples in the Implementation"], [
    ("Measurement", "Generate a pulse signal that represents water movement.", "YF-S201 flow sensor; pipe connection; measured-volume calibration procedure."),
    ("Device", "Sample pulses, calculate interval values, display status, retain pending data, and send telemetry.", "M5Stack CoreS3; GPIO2; one-second sampling; three-second upload interval; local persistence."),
    ("Application", "Validate, store, derive, notify, bill, report, and audit.", "Laravel API; MySQL-compatible database; WaterConsumptionService; BillingStatementService; Sanctum credentials."),
    ("Household workspace", "View only the authenticated household's current and historical information.", "Consumption, billing, notifications, goals, requests, settings, downloadable records."),
    ("Municipal workspace", "Manage registered households and devices and review operational records.", "Device lifecycle, billing configuration, leaks, reports, offline payments, service requests, audit logs.")
], [1.25, 2.3, 2.95])
add_heading(doc, "3.2 Data and Control Flow", 2)
for item in [
    "Water passing through the sensor produces pulses. The CoreS3 samples and converts pulses into an interval flow rate and interval liters using its configured calibration factor.",
    "The device keeps a distinct reading identifier for each interval. If the upload fails, the interval remains in device storage and is sent again with the same identifier after connectivity returns.",
    "A commissioned device sends its telemetry to the Laravel API using a scoped credential. The server checks device identity, lifecycle status, and household assignment before accepting the reading.",
    "The application stores the accepted interval, derives the server-side monthly consumption total, evaluates applicable billing settings and possible-leak conditions, and returns authoritative state for the device display and dashboards.",
    "Household and administrator interfaces retrieve only the records permitted for their authenticated role. Administrators can perform governed lifecycle and configuration actions, which are recorded in audit logs."
]:
    add_numbered(doc, item)
add_heading(doc, "3.3 Device and Sensor Operation", 2)
add_body(doc, "The firmware identifies the panel as M5Stack CoreS3 and samples the sensor at a one-second interval. The nominal upload interval is three seconds. The device uses the initial project calibration of 450 pulses per liter. The value is deliberately configurable because the installed sensor, water pressure, pipe geometry, and real mounting conditions can affect actual measurement. A measured container test must establish a corrected factor before the device is presented as calibrated for field use.")
add_body(doc, "The device screen supports local situational awareness without exposing credentials. It shows measured flow, accumulated context, Wi-Fi state, server state, monitoring status, and server information returned after accepted telemetry. The firmware does not decide municipal thresholds or declare the authoritative monthly total; these decisions remain on the server so a device cannot unilaterally modify household consumption or billing rules.")
add_heading(doc, "3.4 Web Application and Data Services", 2)
add_body(doc, "The web application is implemented with Laravel, Blade-rendered user interfaces, API controllers, services, migrations, and automated tests. Device telemetry is received by API endpoints and processed through service logic that protects household association, supports duplicate reading detection, derives current-period totals, and creates notifications. The data model includes users, devices, water readings, billing settings, billing statements, payments, leak alerts, notifications, service requests, and audit logs.")
add_body(doc, "Billing support is based on the effective configuration that applies to a household and billing period. The implemented statement service combines consumption in cubic meters with the selected rate, fixed monthly charge, and subsidy or deduction; it applies a zero floor to the current charge, carries forward an opening balance, and incorporates confirmed offline payments. These calculations support monitoring and statement generation, but final municipal adoption still requires approval of rates, procedures, reconciliation, and official records.")
add_heading(doc, "3.5 Key Functional Features", 2)
add_simple_table(doc, ["Feature", "Implemented Behavior", "Boundary"], [
    ("Meter commissioning", "Administrator registration, one-time activation, hardware binding, credential issuance, rotation, suspension, and decommissioning.", "Requires authorized municipal action and secure handling of activation information."),
    ("Reliable telemetry", "Pending intervals persist locally; retries keep the same ID; duplicates are recognized server-side.", "Does not guarantee visibility during a prolonged power, network, or server outage."),
    ("Consumption monitoring", "Server-derived monthly total, live status, history, trend views, and offline state indicators.", "Values depend on correctly installed and calibrated sensing hardware."),
    ("Billing information", "Configurable rates, fixed charges, subsidies or deductions, statements, receipts, and offline-payment ledger.", "System outputs remain subject to official municipal billing review."),
    ("Possible-leak workflow", "Configured flow threshold and duration, household and optional administrator notification, alert lifecycle, and follow-up.", "Alert is not a confirmed diagnosis and does not close a water valve."),
    ("Privacy and audit", "Role-scoped records, expiring web sessions, device credentials, audit logs, and redaction of sensitive audit metadata.", "Operational security still requires deployment configuration, backups, and administrator practice.")
], [1.4, 3.25, 1.85])
add_heading(doc, "3.6 Security, Reliability, and Data Integrity", 2)
add_body(doc, "The system uses separate concerns for browser users and commissioned devices. Household and administrator records are scoped to authenticated users, while a physical meter is provisioned through its own lifecycle. Device credentials can be revoked when a meter is replaced or removed from service. Sensitive values such as passwords, access tokens, and activation codes are redacted from audit metadata. These controls reduce avoidable exposure, but secure deployment also requires protected server configuration, credential handling, routine updates, and tested backups.")
add_body(doc, "Reliability is addressed in the telemetry path through local pending-reading persistence and idempotency. A reading that is retransmitted after a failed connection retains its identity, and the server returns the existing accepted record instead of adding the interval twice. The current monthly total is calculated from accepted database records, which prevents a client-supplied cumulative number from replacing the server record. This design improves recovery behavior but does not remove the need to observe the device, network, and database during an actual field pilot.")
add_heading(doc, "3.7 Verification Boundary", 2)
add_body(doc, "The implemented project has automated coverage for device activation and telemetry, invalid and suspended credentials, duplicate-reading handling, monthly consumption rollover, retry acknowledgment, billing calculations, statements and offline payments, household data isolation, alerts, responsive dashboard behavior, and selected security controls. The latest project verification run completed 181 automated tests with 1,178 assertions passing. This result supports the tested software behaviors only. It does not prove physical sensor accuracy, real plumbing performance, connection recovery in the intended installation, backup restoration, or unassisted household usability.")
add_heading(doc, "References", 1)
for reference in [
    "Andrić, I., Vrsalović, A., Perković, T., Aglić Čuvić, A., & Šolić, P. (2022). IoT approach towards smart water usage. Journal of Cleaner Production, 365, 132792. https://doi.org/10.1016/j.jclepro.2022.132792",
    "Laravel. (2026). Laravel Sanctum documentation. https://laravel.com/docs/sanctum",
    "M5Stack. (n.d.). CoreS3 documentation. https://docs.m5stack.com/en/core/CoreS3",
    "Bragalli, C., Neri, M., & Toth, E. (2019). Effectiveness of smart meter-based urban water loss assessment in a real network with synchronous and incomplete readings. Environmental Modelling & Software, 112, 128-142. https://doi.org/10.1016/j.envsoft.2018.10.010",
    "Tasong, A. C., & Abao, R. P. (2019). Design and development of an IoT application with visual analytics for water consumption monitoring. Procedia Computer Science, 157, 205-213. https://doi.org/10.1016/j.procs.2019.08.159"
]:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    p.paragraph_format.space_after = Pt(5)
    p.add_run(reference)

for section in doc.sections:
    for paragraph in section.header.paragraphs:
        paragraph.text = "SMART WATER CONSUMPTION MONITORING SYSTEM"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 0, 0)

doc.save(OUTPUT)
print(OUTPUT)
