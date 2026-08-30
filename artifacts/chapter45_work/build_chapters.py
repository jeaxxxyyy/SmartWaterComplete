from pathlib import Path
from datetime import date

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\laragon\www\SmartWater_Update")
WORK = ROOT / "artifacts" / "chapter45_work"
OUT = ROOT / "artifacts" / "Smart_Water_Monitoring_Consumption_System_Chapters_4_and_5.docx"
WORK.mkdir(parents=True, exist_ok=True)

NAVY = "153B5B"
BLUE = "1E6A8D"
TEAL = "1B8A8F"
PALE = "EAF4F7"
LIGHT = "F5F8FA"
MID = "D5E3EA"
RED = "9E3D3D"
GREEN = "26734D"
TEXT = "1F2933"


def font(size, bold=False):
    paths = [
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibri.ttf",
    ]
    if bold:
        paths = [r"C:\Windows\Fonts\arialbd.ttf", r"C:\Windows\Fonts\calibrib.ttf"]
    for p in paths:
        if Path(p).exists():
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


F_TITLE = font(42, True)
F_HEAD = font(28, True)
F_BODY = font(22)
F_SMALL = font(18)
F_SMALL_B = font(18, True)


def rounded_box(draw, xy, fill, outline=NAVY, radius=18, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill="#" + fill, outline="#" + outline, width=width)


def center_text(draw, xy, text, fnt, fill="#1F2933", gap=5):
    x1, y1, x2, y2 = xy
    lines = text.split("\n")
    boxes = [draw.textbbox((0, 0), line, font=fnt) for line in lines]
    total = sum(b[3] - b[1] for b in boxes) + gap * (len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, box in zip(lines, boxes):
        w = box[2] - box[0]
        h = box[3] - box[1]
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=fnt, fill=fill)
        y += h + gap


def arrow(draw, start, end, fill="#1E6A8D", width=5):
    draw.line([start, end], fill=fill, width=width)
    import math
    ang = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 14
    a1 = ang + 2.6
    a2 = ang - 2.6
    p1 = (end[0] + size * math.cos(a1), end[1] + size * math.sin(a1))
    p2 = (end[0] + size * math.cos(a2), end[1] + size * math.sin(a2))
    draw.polygon([end, p1, p2], fill=fill)


def make_architecture(path):
    im = Image.new("RGB", (1500, 820), "white")
    d = ImageDraw.Draw(im)
    d.text((55, 35), "Smart Water System Architecture", font=F_TITLE, fill="#153B5B")
    boxes = [
        ((70, 180, 350, 390), "Flow Sensor +\nM5Stack CoreS3", "EAF4F7"),
        ((455, 180, 745, 390), "Wi-Fi / HTTPS\nJSON Telemetry", "F5F8FA"),
        ((850, 130, 1160, 440), "Laravel 13 API\nSanctum Auth\nBusiness Services", "EAF4F7"),
        ((1240, 180, 1450, 390), "SQLite / MySQL\nCore Records", "F5F8FA"),
        ((350, 560, 670, 740), "Household Workspace\nLive usage, billing, alerts", "EAF4F7"),
        ((825, 560, 1165, 740), "Administrator Workspace\nAccounts, meters, reports", "EAF4F7"),
    ]
    for xy, label, fill in boxes:
        rounded_box(d, xy, fill)
        center_text(d, xy, label, F_BODY)
    arrow(d, (350, 285), (455, 285))
    arrow(d, (745, 285), (850, 285))
    arrow(d, (1160, 285), (1240, 285))
    arrow(d, (1000, 440), (980, 560))
    arrow(d, (930, 440), (650, 560))
    im.save(path)


def make_workflow(path):
    im = Image.new("RGB", (1500, 880), "white")
    d = ImageDraw.Draw(im)
    d.text((55, 30), "Commissioning and Telemetry Workflow", font=F_TITLE, fill="#153B5B")
    items = [
        ("1", "Administrator registers meter", "Asset number, serial, household,\ninstallation and lifecycle details"),
        ("2", "Device activates once", "UUID + expiring code + hardware ID\nexchange for scoped token"),
        ("3", "Sensor measures interval use", "Pulses converted to liters and\nflow rate; state persisted locally"),
        ("4", "API validates and stores", "Device identity owns the reading;\nduplicate reading_id is idempotent"),
        ("5", "Services calculate results", "Consumption totals, estimated bill,\nthreshold and leak evaluation"),
        ("6", "Users review and act", "Households monitor usage; admins\nmanage devices, alerts and reports"),
    ]
    y = 125
    for i, (n, title, sub) in enumerate(items):
        xy = (170, y, 1330, y + 105)
        rounded_box(d, xy, "EAF4F7" if i % 2 == 0 else "F5F8FA", radius=14, width=2)
        d.ellipse((80, y + 20, 145, y + 85), fill="#1E6A8D")
        center_text(d, (80, y + 20, 145, y + 85), n, F_HEAD, fill="white")
        d.text((205, y + 16), title, font=F_SMALL_B, fill="#153B5B")
        d.multiline_text((205, y + 47), sub, font=F_SMALL, fill="#1F2933", spacing=4)
        if i < len(items) - 1:
            arrow(d, (750, y + 105), (750, y + 125), width=4)
        y += 120
    im.save(path)


def make_erd(path):
    im = Image.new("RGB", (1600, 1000), "white")
    d = ImageDraw.Draw(im)
    d.text((55, 30), "Core Entity Relationship Overview", font=F_TITLE, fill="#153B5B")
    entities = {
        "USERS": (70, 160, 390, 390, ["id (PK)", "role / account_number", "account status", "name / contact details"]),
        "DEVICES": (640, 120, 980, 370, ["id (PK)", "household_id (FK)", "uuid / asset / serial", "hardware / lifecycle status"]),
        "WATER_READINGS": (1210, 160, 1540, 430, ["id (PK)", "user_id, device_id (FK)", "reading_id", "flow / liters / m3", "leak and threshold flags"]),
        "BILLINGS": (70, 610, 390, 860, ["id (PK)", "user_id (FK, optional)", "rate / fixed charge", "threshold and leak rules", "effective date / status"]),
        "BILLING_STATEMENTS": (570, 640, 1010, 920, ["id (PK)", "user_id (FK)", "billing month", "usage and charge snapshot", "payment status / balance"]),
        "LEAK_ALERTS": (1210, 650, 1540, 920, ["id (PK)", "user_id (FK)", "reading references", "status / resolution", "notification reference"]),
    }
    for name, (x1, y1, x2, y2, fields) in entities.items():
        rounded_box(d, (x1, y1, x2, y2), "F5F8FA", radius=12, width=3)
        d.rectangle((x1, y1, x2, y1 + 55), fill="#153B5B")
        center_text(d, (x1, y1, x2, y1 + 55), name, F_SMALL_B, fill="white")
        yy = y1 + 75
        for field in fields:
            d.text((x1 + 18, yy), field, font=F_SMALL, fill="#1F2933")
            yy += 38
    arrow(d, (390, 265), (640, 245))
    d.text((465, 220), "1 to many", font=F_SMALL, fill="#1E6A8D")
    arrow(d, (980, 255), (1210, 285))
    d.text((1030, 220), "1 to many", font=F_SMALL, fill="#1E6A8D")
    arrow(d, (250, 390), (250, 610))
    arrow(d, (390, 335), (700, 640))
    arrow(d, (850, 370), (850, 640))
    arrow(d, (1375, 430), (1375, 650))
    d.text((1120, 540), "readings may trigger", font=F_SMALL, fill="#1E6A8D")
    im.save(path)


def make_interface_map(path, admin=False):
    im = Image.new("RGB", (1600, 950), "#F5F8FA")
    d = ImageDraw.Draw(im)
    title = "Administrator Operations Workspace" if admin else "Household Monitoring Workspace"
    d.rectangle((0, 0, 1600, 95), fill="#153B5B")
    d.text((45, 24), title, font=F_HEAD, fill="white")
    d.rectangle((0, 95, 330, 950), fill="#EAF4F7")
    nav = (["Overview", "Customer Management", "Device Management", "Water Rate Management", "Billing", "Consumption Monitoring", "Leak Monitoring", "Reports", "Notifications"] if admin else ["Dashboard", "Live Monitoring", "Consumption History", "Estimated Billing", "Leak Detection", "Notifications", "Account Settings"])
    y = 140
    for i, item in enumerate(nav):
        fill = "#1E6A8D" if i == 0 else "#EAF4F7"
        color = "white" if i == 0 else "#153B5B"
        d.rounded_rectangle((25, y, 305, y + 58), radius=10, fill=fill)
        d.text((48, y + 17), item, font=F_SMALL, fill=color)
        y += 72
    if admin:
        cards = ["Active households", "Online meters", "Open leak alerts", "Unread notices"]
        panels = ["12-month consumption trend", "Operational follow-ups", "Recent audit and device events"]
    else:
        cards = ["Current flow rate", "Estimated bill", "Leak status", "Meter connection", "Notifications"]
        panels = ["Usage comparison and meter history", "Latest meter activity", "Recent alerts and updates"]
    x = 375
    width = 220 if len(cards) == 5 else 270
    for item in cards:
        rounded_box(d, (x, 135, x + width, 280), "FFFFFF", outline="D5E3EA", radius=12, width=2)
        d.text((x + 18, 160), item, font=F_SMALL_B, fill="#153B5B")
        d.text((x + 18, 215), "--", font=F_HEAD, fill="#1E6A8D")
        x += width + 25
    y = 335
    heights = [245, 140, 140]
    for item, h in zip(panels, heights):
        rounded_box(d, (375, y, 1540, y + h), "FFFFFF", outline="D5E3EA", radius=12, width=2)
        d.text((405, y + 24), item, font=F_SMALL_B, fill="#153B5B")
        if h > 200:
            pts = [(430, y + 175), (620, y + 130), (790, y + 160), (970, y + 85), (1160, y + 115), (1450, y + 65)]
            d.line(pts, fill="#1B8A8F", width=6)
        else:
            d.line((410, y + 75, 1495, y + 75), fill="#D5E3EA", width=3)
            d.line((410, y + 103, 1380, y + 103), fill="#EAF4F7", width=3)
        y += h + 28
    im.save(path)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    trPr.append(cant_split)


def set_column_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tcPr = row.cells[idx]._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(width * 1440)))
            tcW.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for idx, h in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_with_next = True
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(font_size)
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if ridx % 2 == 1:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)
    if widths:
        set_column_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(10)


def add_figure(doc, path, caption, width=6.45):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption.split(".", 1)[0])
    add_caption(doc, caption)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.add_run(item)


def add_note(doc, title, text, fill=PALE):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.style = "Table Grid"
    set_repeat_table_header(t.rows[0])
    prevent_row_split(t.rows[0])
    set_column_widths(t, [6.5])
    cell = t.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 130, 160, 130, 160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_chapter_title(doc, number, title):
    if len(doc.paragraphs) > 1:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"CHAPTER {number}")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(20)
    r2 = p2.add_run(title.upper())
    r2.bold = True
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(14)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Inches(0.5)

    for style_name, size, color, before, after in [
        ("Heading 1", 13, NAVY, 12, 6),
        ("Heading 2", 12, BLUE, 10, 4),
        ("Heading 3", 11, TEAL, 8, 3),
    ]:
        st = styles[style_name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.first_line_indent = Inches(0)

    for sec in doc.sections:
        header = sec.header
        hp = header.paragraphs[0]
        hp.text = "SMART WATER MONITORING CONSUMPTION SYSTEM"
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in hp.runs:
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor.from_string(BLUE)
        footer = sec.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run("Page ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        run._r.addnext(fld)
        for r in fp.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9)


def build():
    arch = WORK / "architecture.png"
    workflow = WORK / "workflow.png"
    erd = WORK / "erd.png"
    household = WORK / "household_interface_map.png"
    admin = WORK / "admin_interface_map.png"
    make_architecture(arch)
    make_workflow(workflow)
    make_erd(erd)
    make_interface_map(household, False)
    make_interface_map(admin, True)

    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(115)
    r = p.add_run("SMART WATER MONITORING\nCONSUMPTION SYSTEM")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(25)
    p2.add_run("CHAPTERS IV AND V").bold = True
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(120)
    p3.add_run("Project-based manuscript draft\nCalape, Bohol").italic = True
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(100)
    p4.add_run("Prepared from the implemented Laravel and ESP32 project\nAugust 2026")

    add_chapter_title(doc, "IV", "Methodology, System Design, and Implementation")
    add_body(doc, "This chapter presents the development methodology, requirements, design, database structure, implementation process, and verification results of the Smart Water Monitoring Consumption System. The organization follows the academic flow of the supplied reference document, while all technical descriptions in this chapter are derived from the actual Smart Water project files, routes, database migrations, firmware, interface templates, services, and automated tests.")

    doc.add_heading("4.1 Requirements Analysis and Requirements Documentation", level=1)
    add_body(doc, "Requirements analysis identified three principal actors: the municipal administrator, the household account holder, and the commissioned smart meter. The administrator manages household records, municipal meters, water-rate configuration, billing statements, leak incidents, notifications, analytics, reports, and audit information. The household user reviews live and historical consumption, estimated billing, leak status, notifications, and account settings. The smart meter measures flow-sensor pulses and submits authenticated interval readings to the server.")

    doc.add_heading("4.1.1 Functional Requirements", level=2)
    functional = [
        ("FR-01", "User and access management", "Administrator creates and manages household accounts; public registration is restricted; account status governs access."),
        ("FR-02", "Municipal device lifecycle", "Register, assign, activate, rotate credentials, reassign, suspend, and decommission meters."),
        ("FR-03", "Telemetry ingestion", "Accept authenticated interval readings containing a stable reading identifier, flow rate, liters used, meter state, and device diagnostics."),
        ("FR-04", "Consumption monitoring", "Display live readings, daily/weekly/monthly history, summaries, comparisons, and availability states."),
        ("FR-05", "Estimated billing", "Apply the latest applicable water rate, fixed charge, subsidy, and billable consumption to estimates and monthly statements."),
        ("FR-06", "Leak and threshold detection", "Evaluate server-configured thresholds, notify affected users, and support an administrator resolution workflow."),
        ("FR-07", "Notifications", "Create, read, archive, retry, and remove household or administrative notices according to role and ownership."),
        ("FR-08", "Reports and exports", "Produce consumption, billing, leak, notification, and operational reports with filtering and export support."),
        ("FR-09", "Audit and analytics", "Record administrative/security activity and present household- or administrator-scoped analytics."),
        ("FR-10", "Account maintenance", "Allow authenticated users to update permitted profile fields and change passwords."),
    ]
    add_table(doc, ["ID", "Requirement", "Verified Project Behavior"], functional, [0.65, 1.75, 4.10], 9)

    doc.add_heading("4.1.2 Non-Functional Requirements", level=2)
    nonfunctional = [
        ("Security", "Laravel Sanctum tokens, role middleware, scoped device credentials, hashed one-time activation codes, token revocation, and audit redaction."),
        ("Reliability", "Stable reading identifiers prevent double counting; unsent meter intervals are retained in ESP32 Preferences until acknowledged."),
        ("Performance", "The meter uploads at three-second intervals and the web interface is designed to refresh live information at the same cadence."),
        ("Data integrity", "Foreign keys, validation, unique identifiers, server-owned household/device associations, precision-preserving numeric fields, and historical charge snapshots."),
        ("Usability", "Responsive administrator and household workspaces, accessible labels, keyboard-oriented controls, role identification, and light/dark theme support."),
        ("Maintainability", "Controllers are grouped by feature; models represent domain entities; services isolate billing, analytics, consumption, and audit logic."),
        ("Portability", "The web application supports SQLite or MySQL, while the field device targets ESP32-S3/M5Stack CoreS3 hardware."),
    ]
    add_table(doc, ["Quality Attribute", "System Response"], nonfunctional, [1.45, 5.05], 9.2)

    doc.add_heading("4.1.3 Use-Case Description", level=2)
    use_cases = [
        ("Manage household accounts", "Administrator", "Authenticated administrator", "Create, search, filter, edit, suspend, deactivate, or archive a household account.", "A role-scoped household record with a standardized account number."),
        ("Commission a smart meter", "Administrator + Device", "Registered and eligible household", "Register the asset, issue a one-time code, activate from the ESP32, and bind the hardware identity.", "Active scoped device credential and assigned meter."),
        ("Upload a reading", "Smart meter", "Active assigned device and valid token", "Measure pulses, calculate interval values, submit JSON, validate identity, and save once.", "Owned water reading and updated server total."),
        ("Review consumption", "Household", "Authenticated active account", "Open live monitoring or history; select a period; retrieve scoped readings and summaries.", "Current flow, usage totals, trends, and availability state."),
        ("Generate billing statement", "Administrator", "Valid rate and billable readings", "Select the billing month, aggregate consumption, apply charge rules, and save a snapshot.", "Finalized statement with status and balances."),
        ("Resolve a leak alert", "Administrator", "Leak incident exists", "Review household and reading evidence, notify, classify, resolve, mark false alarm, or reopen.", "Updated incident state and traceable response."),
    ]
    add_table(doc, ["Use Case", "Actor", "Precondition", "Main Flow", "Result"], use_cases, [1.15, 0.90, 1.15, 2.05, 1.25], 8.2)

    doc.add_heading("4.2 System Design, Software, and Processes", level=1)
    add_body(doc, "The system uses an Internet of Things client-server design. The physical meter performs measurement and short-term persistence, while the Laravel server remains authoritative for identity, ownership, cumulative usage, billing rules, threshold evaluation, and record storage. The browser interfaces consume protected JSON endpoints and present role-specific workspaces.")
    add_figure(doc, arch, "Figure 4.1. Smart Water Monitoring Consumption System architecture.")

    doc.add_heading("4.2.1 Telemetry and Billing Process", level=2)
    add_body(doc, "Commissioning begins in the administrator workspace. The device exchanges its UUID, expiring activation code, hardware ID, and firmware version for a scoped credential. Each reading includes a stable reading_id derived from the registered device identity and local sequence. If transmission fails, the pending interval remains stored in the device. When the API acknowledges the reading, the local sequence advances and the interval accumulator is cleared. The server validates ownership and computes authoritative totals before exposing information to users.")
    add_figure(doc, workflow, "Figure 4.2. Device commissioning, telemetry, calculation, and user-response workflow.")

    doc.add_heading("4.2.2 Technology Used", level=2)
    tech = [
        ("Backend", "PHP 8.3+; Laravel Framework 13.8; Laravel Sanctum 4.0", "Routing, validation, authentication, API controllers, services, reports, and persistence."),
        ("Frontend", "Blade, JavaScript, Tailwind CSS 4, Vite 8", "Responsive household and administrator workspaces, charts, filters, forms, and live refresh."),
        ("Database", "SQLite or MySQL", "Users, meters, readings, rates, billing statements, notifications, leak alerts, audit logs, sessions, and tokens."),
        ("Embedded device", "ESP32-S3 / M5Stack CoreS3; M5Unified; Wi-Fi; HTTPClient; Preferences", "Pulse measurement, device display, activation, reliable telemetry, and local pending-state storage."),
        ("Reporting", "Dompdf 3.1.2 and CSV export endpoints", "Printable operational reports and downloadable tabular records."),
        ("Testing", "PHPUnit 12.5.12 through Laravel test runner", "Feature verification for access control, telemetry, billing, alerts, analytics, devices, and interfaces."),
    ]
    add_table(doc, ["Layer", "Technology", "Purpose"], tech, [1.05, 2.55, 2.90], 8.8)

    doc.add_heading("4.3 Entity Relationship Diagram", level=1)
    add_body(doc, "The principal database relationships connect household users to their assigned devices, readings, statements, notifications, and leak incidents. A device may be temporarily unassigned during municipal inventory or maintenance, but an uploaded reading is tied to both the authenticated device and its current household. Billing configuration can be municipal or user-specific, and finalized statements preserve the charge values applied at generation time.")
    add_figure(doc, erd, "Figure 4.3. Core entity relationship overview (operational fields summarized).", 6.5)

    doc.add_heading("4.4 Data Dictionary", level=1)
    add_body(doc, "The following data dictionary summarizes the operational tables most directly involved in monitoring, billing, alerting, and administration. Framework-support tables for cache, jobs, sessions, password resets, and personal access tokens remain part of the implementation but are omitted from the main discussion for clarity.")
    dictionaries = [
        ("users", "id; name; email; role; account_number; address; phone; is_active; account_status; created_by_admin; timestamps", "Household and administrator identities, contact details, and access state."),
        ("devices", "id; household_id; device_uuid; asset_number; manufacturer_serial; hardware_id; firmware_version; installation_date; status; last_seen_at; activation metadata", "Municipal meter registry, assignment, identity binding, health, and lifecycle."),
        ("water_readings", "id; user_id; device_id; reading_id; flow_rate; liters_used; cubic_meter; estimated_bill; rate; threshold and leak flags; connectivity fields; recorded_at", "Time-stamped, device-owned interval telemetry and derived monitoring state."),
        ("billings", "id; user_id; water_rate; fixed_monthly_charge; subsidy_deduction; currency; effective_date; threshold_liters; leak rules; is_active", "Effective-dated municipal or household billing and monitoring configuration."),
        ("billing_statements", "id; user_id; billing_month; opening_balance; total_liters; total_cubic_meter; water_rate; fixed charge; subsidy; current_charge; amount_paid; remaining_balance; status", "Historical monthly statement and payment-state snapshot."),
        ("notifications", "id; user_id; created_by; type; title; message; read state; batch; delivery status/error; sent/scheduled/archived timestamps", "Household notices, administrator announcements, delivery lifecycle, and read status."),
        ("leak_alerts", "id; user_id; first/latest reading IDs; notification_id; status; detected/resolved timestamps; resolution notes; last notified date", "Operational incident record supporting notification, classification, resolution, false-alarm handling, and reopening."),
        ("audit_logs", "id; user_id; action; entity type/id; metadata; IP address; user agent; request ID; created_at", "Trace of administrative and security-sensitive activity with contextual metadata."),
    ]
    add_table(doc, ["Table", "Important Fields", "Purpose"], dictionaries, [1.15, 3.35, 2.00], 8.2)

    doc.add_heading("4.5 Development and Testing", level=1)
    add_body(doc, "Development is organized around feature-specific API controllers and domain services. WaterConsumptionService and HouseholdConsumptionService handle aggregation and availability semantics; BillingStatementService preserves monthly charge logic; AnalyticsService produces range-based summaries; and the embedded firmware handles acquisition, display, activation, and delivery reliability. Verification combines automated feature tests with a production frontend build.")
    test_rows = [
        ("Laravel automated test suite", "138 tests; 885 assertions", "137 passed; 1 failed", "18.017 seconds", "The failing case expected one record but received two in an administrator consumption-history response."),
        ("Frontend production build", "Vite 8.0.16; 4 modules transformed", "Passed", "6.75 seconds", "CSS, boot JavaScript, main JavaScript, and manifest were produced successfully."),
    ]
    add_table(doc, ["Verification", "Coverage / Output", "Status", "Duration", "Observation"], test_rows, [1.25, 1.40, 0.70, 0.75, 2.40], 8.2)
    add_note(doc, "Current verification status", "The project is substantially covered by automated tests, but the suite is not fully green as of 24 August 2026. One consumption-monitoring test must be investigated before claiming complete regression-test acceptance.", "FCEBEC")

    doc.add_heading("4.5.1 Major Test Areas", level=2)
    add_bullets(doc, [
        "Authentication and account restrictions, including blocked, suspended, archived, deleted, and unauthorized users.",
        "Municipal device registration, activation, hardware binding, credential rotation, reassignment, health, and decommissioning.",
        "Telemetry validation, device ownership, idempotent reading identifiers, server-authoritative totals, and precision.",
        "Household live readings, time-zone-aware periods, cumulative-reset handling, missing-data semantics, and billing estimates.",
        "Water-rate configuration, fixed monthly charges, statement generation, status changes, and historical billing fields.",
        "Leak alert creation, notification, resolution, false-alarm handling, reopening, and role protection.",
        "Administrator reporting, analytics, customer management, exports, notification lifecycle, and interface accessibility contracts.",
    ])

    doc.add_heading("4.6 Implementation Plan", level=1)
    plan = [
        ("1. Environment preparation", "Municipal IT / Developers", "Configure PHP, database, application key, mail/queue behavior, timezone, LAN URL, and production assets.", "Environment checklist and successful health endpoint."),
        ("2. Database deployment", "Developers", "Back up the target database, run migrations and controlled seeders, and verify indexes and foreign keys.", "Current schema and verified baseline records."),
        ("3. Administrator setup", "System administrator", "Create authorized staff accounts, set water rates and alert rules, and verify audit logging.", "Operational administrative workspace."),
        ("4. Household onboarding", "Authorized staff", "Create household records with required address, contact details, role, status, and account number.", "Active household accounts ready for assignment."),
        ("5. Meter commissioning", "Technician + Administrator", "Register each meter, record asset details, assign household, activate once, and verify last-seen health.", "Bound device and successful telemetry."),
        ("6. Pilot verification", "Project team", "Compare measured volume with a calibrated reference, verify network recovery, review alerts, and reconcile billing estimates.", "Signed pilot acceptance record."),
        ("7. Production rollout", "Municipality", "Deploy in phases, train staff and households, monitor audit/leak/device dashboards, and maintain backups.", "Controlled operational adoption."),
    ]
    add_table(doc, ["Phase", "Owner", "Activity", "Expected Output"], plan, [1.15, 1.15, 2.85, 1.35], 8.4)

    doc.add_heading("4.6.1 Implementation Issues and Challenges", level=2)
    challenges = [
        ("Sensor calibration", "Pulse-to-liter conversion varies by sensor, pipe conditions, pressure, and installation.", "Perform volumetric calibration per installed meter and retain calibration records."),
        ("Network interruption", "LAN or Wi-Fi loss can delay near-real-time visibility.", "Use the firmware pending-reading state, stable IDs, health windows, and field connectivity checks."),
        ("Duplicate or noisy pulses", "Electrical noise may inflate flow readings.", "Use minimum pulse intervals, abnormal-burst rejection, and site validation."),
        ("Credential handling", "Activation codes or tokens may be exposed during commissioning.", "Issue short-lived single-use codes, remove codes after activation, rotate credentials, and restrict administrator access."),
        ("Billing configuration", "Incorrect effective dates or rate settings can affect estimates and statements.", "Require authorized review, retain charge snapshots, and reconcile pilot statements before rollout."),
        ("Test-suite regression", "One current consumption-history test has a record-count mismatch.", "Isolate the query/date fixture, correct the root cause, and rerun the complete suite before acceptance."),
    ]
    add_table(doc, ["Issue", "Effect", "Mitigation"], challenges, [1.30, 2.20, 3.00], 8.7)

    doc.add_heading("4.7 System Implementation and Results", level=1)
    add_body(doc, "The implemented application provides a complete path from a registered physical meter to role-scoped monitoring and administrative action. The meter reports measured intervals; the API validates device identity; the database preserves raw and derived values; services calculate consumption and estimated billing; and the user interfaces expose monitoring, leak, notification, billing, analytics, and reporting functions. The successful production asset build and 137 passing automated tests demonstrate broad functional coverage, while the remaining failing test identifies a specific acceptance item rather than an undocumented weakness.")
    add_note(doc, "Interpretation", "The implementation is suitable for continued pilot testing and controlled demonstration. Production acceptance should follow calibration, backup, security, network-recovery, and full regression-test sign-off.")

    add_chapter_title(doc, "V", "Presentation, Analysis, and Interpretation of Results")
    add_body(doc, "This chapter presents the implemented interfaces and interprets the results of the Smart Water Monitoring Consumption System against its operational objectives. The figures are structured interface maps derived from the current Blade templates and implemented navigation. They summarize the actual workspace composition without presenting fabricated household consumption values.")

    doc.add_heading("5.1 Landing and Authentication Interface", level=1)
    add_body(doc, "The public entry points provide a landing page, splash page, and login page. Authentication is API-based and uses Laravel Sanctum. Public registration is intentionally unavailable; household accounts are created by authorized administrators. Successful authentication directs the user to the workspace corresponding to the account role. The login interface includes an accessible password visibility control, while the shared page shell provides theme controls and responsive navigation.")
    add_table(doc, ["Interface Element", "Function", "Interpretation"], [
        ("Landing page", "Introduces the water-monitoring service and directs users to authentication.", "Keeps public content separate from protected operational data."),
        ("Login form", "Validates email and password and requests a scoped session token.", "Centralizes access and blocks inactive, suspended, archived, incomplete, or unauthorized accounts."),
        ("Role redirect", "Routes administrators and households to their own workspace.", "Reduces accidental exposure of administrative controls."),
    ], [1.35, 2.55, 2.60], 9)

    doc.add_heading("5.2 Household Dashboard Interface", level=1)
    add_body(doc, "The household workspace uses a compact authenticated layout with a collapsible sidebar and recent-notification control. The dashboard summarizes current flow rate, estimated bill, leak status, meter connection, and notification count. It also provides usage comparison, meter history, latest meter activity, and recent alerts. Range controls support predefined periods and a validated custom date range.")
    add_figure(doc, household, "Figure 5.1. Household workspace interface map derived from the implemented dashboard template.")

    doc.add_heading("5.2.1 Live Monitoring", level=2)
    add_body(doc, "Live Monitoring displays recent readings from the M5Stack device and a real-time consumption trend. Date filters and selectable row counts allow the household to narrow the displayed interval. The live state is based on the meter heartbeat rather than on a nonzero flow rate, so a connected meter can correctly report zero water use. Missing readings are not silently interpreted as zero consumption.")

    doc.add_heading("5.2.2 Consumption History", level=2)
    add_body(doc, "Consumption History groups information by daily, weekly, and monthly periods and provides summary cards, a trend chart, and detailed reading totals. The implemented consumption service distinguishes interval readings from cumulative device totals, handles safe differences when cumulative meters reset, and uses the configured application timezone when constructing reporting periods.")

    doc.add_heading("5.2.3 Estimated Billing", level=2)
    add_body(doc, "The billing interface presents a single current estimate, a transparent billing breakdown, assigned-meter information, and finalized statement history. The calculation requires valid billable readings and an applicable active rate. Current charges may include the configured rate per cubic meter, fixed monthly charge, subsidy deduction, and prior balance according to the statement logic. When inputs are unavailable, the interface presents a specific availability message rather than a misleading zero estimate.")

    doc.add_heading("5.2.4 Leak Detection and Notifications", level=2)
    add_body(doc, "Leak Detection presents the latest meter signal, configured threshold, and response status. The notification inbox contains account, billing, consumption, meter, and leak-related notices. Received notifications may be marked as read or removed from the household view without erasing the municipality's source record. This separation protects operational traceability while allowing households to manage their inbox.")

    doc.add_heading("5.2.5 Account Settings", level=2)
    add_body(doc, "Account Settings allows the household to update permitted contact information and change the password. Protected fields—including role, administrative creation state, and server-owned account controls—remain outside normal household editing. This supports self-service maintenance without weakening authorization boundaries.")

    doc.add_heading("5.3 Administrator Operations Interface", level=1)
    add_body(doc, "The administrator workspace consolidates municipal operations in one responsive shell. Its overview includes account and meter health, consumption trend, operational follow-ups, notifications, and recent administrative or device events. Global search and feature-specific filters support efficient retrieval of households, meters, billing records, incidents, reports, and notices.")
    add_figure(doc, admin, "Figure 5.2. Administrator workspace interface map derived from the implemented admin template.")

    admin_modules = [
        ("Customer Management", "Create household accounts; search by name, email, account, address, or meter; filter status; update, suspend, deactivate, archive, or unassign a meter.", "Administrative account creation and explicit status controls support controlled onboarding."),
        ("Device Management", "Register municipal assets, assign households, issue activation codes, review hardware/firmware/last-seen health, rotate credentials, reassign, and decommission.", "The registry separates equipment inventory from interactive user accounts."),
        ("Water Rate Management", "Configure rates, fixed charges, currency, effective dates, consumption thresholds, leak thresholds, durations, and notification behavior.", "Effective-dated configuration supports transparent billing and monitoring rules."),
        ("Billing Management", "Search estimated and finalized records, filter period/type/status, generate statements, review calculation details, and mark payments.", "A preserved statement snapshot improves historical consistency."),
        ("Consumption Monitoring", "Search households, filter readings, open daily/weekly/monthly history, and inspect meter state and totals.", "Administrators can investigate usage without weakening household data scoping."),
        ("Leak Monitoring", "Review affected household, latest flow, notice history, and resolution state; resolve, classify false alarm, or reopen.", "Leak events become managed incidents instead of one-time flags."),
        ("Reports", "Preview and export consumption, billing, notification, and leak reports for a selected range.", "Operational data is available for municipal review and record keeping."),
        ("Notification Management", "Compose official notices, select audience and category, track delivery/read status, retry, archive, and delete sent records.", "Delivery lifecycle and audit information support accountable communication."),
    ]
    add_table(doc, ["Module", "Implemented Functions", "Result / Interpretation"], admin_modules, [1.30, 3.25, 1.95], 8.1)

    doc.add_heading("5.4 Smart Meter Interface and Field Operation", level=1)
    add_body(doc, "The M5Stack CoreS3 screen presents Wi-Fi status, server connection, flow rate, total liters, cubic-meter usage, and leak indication. Firmware samples pulse counts every second and attempts an upload every three seconds. It rejects very short duplicate edges and abnormal pulse bursts, retains pending readings across failures, and clears the interval only after server acknowledgement. These behaviors reduce the risk of lost or double-counted usage during temporary connectivity problems.")
    add_table(doc, ["Device Indicator", "Source", "Operational Meaning"], [
        ("Wi-Fi status", "ESP32 connection state", "Whether the meter is connected to the configured local network."),
        ("Server status", "Latest telemetry response", "Whether the API acknowledged a recent upload."),
        ("Flow rate", "Pulse count over elapsed time", "Current water-flow estimate in liters per minute."),
        ("Total liters", "Server total when available; otherwise device total", "Running consumption display with server authority preferred."),
        ("Usage in cubic meters", "Liter conversion", "Consumption unit used by water-rate calculations."),
        ("Leak indication", "Server response", "Whether configured leak logic marked the current event as possible leakage."),
    ], [1.35, 2.05, 3.10], 8.8)

    doc.add_heading("5.5 Analysis of System Results", level=1)
    results = [
        ("Near-real-time monitoring", "Three-second device upload and dashboard refresh design; precise interval fields; scoped live endpoints.", "Implemented", "Supports timely household and municipal visibility, subject to network availability."),
        ("Reliable device identity", "Municipal registry, one-time activation, hardware binding, scoped token, and lifecycle revocation.", "Implemented", "Reduces unauthorized or misassigned telemetry."),
        ("Accurate record ownership", "Server derives household and device from the authenticated meter and rejects invalid associations.", "Implemented", "Prevents clients from selecting another household for a reading."),
        ("Billing transparency", "Effective rate selection, fixed charge, subsidy, statements, explanation states, and historical charge snapshot.", "Implemented", "Provides traceable estimates and finalized monthly records."),
        ("Leak response", "Threshold evaluation, once-per-event notice behavior, incident list, follow-up, resolution, false alarm, and reopening.", "Implemented", "Transforms sensor conditions into an actionable workflow."),
        ("Regression acceptance", "137 of 138 automated tests passed; frontend production build passed.", "Conditional", "Broad coverage is present, but one consumption-history mismatch remains before full acceptance."),
    ]
    add_table(doc, ["Objective", "Evidence from Implementation", "Status", "Interpretation"], results, [1.35, 2.85, 0.80, 1.50], 8.1)

    doc.add_heading("5.6 Findings", level=1)
    add_bullets(doc, [
        "The project implements the full information path from a commissioned flow meter to household and municipal dashboards.",
        "Security boundaries are explicit: administrators control household creation and municipal devices, households receive only their scoped records, and devices hold limited upload credentials.",
        "Data-quality handling is stronger than simple cumulative summation because stable reading identifiers, interval usage, reset-aware aggregation, and missing-data states are represented separately.",
        "Billing results are explainable because current estimates, active rates, fixed charges, subsidies, historical statements, and unavailable-data messages are distinct.",
        "Leak monitoring is operational rather than purely visual because alerts can be notified, followed up, resolved, marked false, and reopened.",
        "The current codebase is close to automated-test acceptance but cannot yet be described as fully passing due to one administrator consumption-history assertion failure.",
    ])

    doc.add_heading("5.7 Conclusions", level=1)
    add_body(doc, "Based on the implemented project and the verification performed, the Smart Water Monitoring Consumption System meets its principal functional purpose: it captures water-flow telemetry from an ESP32-S3/M5Stack CoreS3 meter, stores device- and household-owned readings, presents near-real-time and historical consumption, produces estimated and finalized billing information, manages leak and threshold notifications, and provides municipal tools for accounts, devices, incidents, reports, and audit review. The separation of household, administrator, and device roles creates a sound access-control foundation for municipal deployment.")
    add_body(doc, "The system also demonstrates practical reliability measures through persistent pending readings, idempotent reading identifiers, server-authoritative totals, effective-dated billing rules, and explicit missing-data states. However, software acceptance alone is insufficient for field deployment. Sensor calibration, network recovery, physical installation, backup restoration, security review, and reconciliation with official billing procedures require documented pilot evidence. In addition, the remaining automated-test failure must be resolved before a final claim of complete regression readiness.")

    doc.add_heading("5.8 Recommendations", level=1)
    recommendations = [
        ("Resolve the failing regression test", "Investigate why the administrator consumption-history response returns two records when the fixture expects one; confirm date boundaries, filters, and test isolation; then rerun all 138 tests."),
        ("Conduct calibrated field trials", "Compare measured volume against a certified reference container or meter across low, normal, and high flow, recording calibration constants and acceptable error."),
        ("Formalize acceptance criteria", "Define measurable thresholds for telemetry delay, recovery after outage, billing variance, leak-detection sensitivity, false alarms, availability, and user task completion."),
        ("Strengthen production operations", "Use HTTPS, secure secrets, scheduled backups, restoration drills, log retention, least-privilege accounts, patch management, and network monitoring."),
        ("Validate official billing integration", "Review rates, fixed charges, subsidies, rounding, arrears, due dates, adjustments, and payment posting with the responsible municipal office."),
        ("Train users by role", "Prepare short administrator, technician, and household guides covering onboarding, meter health, alert response, report generation, and account safety."),
        ("Extend notification channels carefully", "After pilot validation, consider queued email or SMS delivery with consent, retry limits, delivery logs, and privacy controls."),
        ("Evaluate long-term analytics", "After sufficient production data exists, review seasonal consumption, meter anomalies, leak recurrence, and conservation outcomes without inferring missing readings as zero."),
    ]
    add_table(doc, ["Recommendation", "Required Action"], recommendations, [1.75, 4.75], 9)

    doc.add_heading("5.9 Chapter Summary", level=1)
    add_body(doc, "Chapter V presented the user and device interfaces, analyzed the implemented modules, mapped the system results to core objectives, and identified the conditions required for deployment. The evidence supports continued pilot use and finalization, with full production acceptance dependent on calibrated field validation, operational controls, official billing review, and resolution of the remaining automated-test failure.")

    props = doc.core_properties
    props.title = "Smart Water Monitoring Consumption System - Chapters IV and V"
    props.subject = "Methodology, implementation, results, conclusions, and recommendations"
    props.author = "Smart Water Monitoring Consumption System Project Team"
    props.keywords = "smart water, monitoring, consumption, IoT, Laravel, ESP32, billing, leak detection"

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
