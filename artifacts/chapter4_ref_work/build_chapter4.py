from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\laragon\www\SmartWater_Update")
SOURCE = ROOT / "artifacts" / "Smart_Water_Monitoring_Consumption_System_Chapters_4_and_5.docx"
WORK = ROOT / "artifacts" / "chapter4_ref_work"
OUT = ROOT / "artifacts" / "Smart_Water_Monitoring_Consumption_System_Chapter_4.docx"
USE_CASE = WORK / "smart_water_use_case.png"
HOUSEHOLD = ROOT / "artifacts" / "chapter45_work" / "household_interface_map.png"
ADMIN = ROOT / "artifacts" / "chapter45_work" / "admin_interface_map.png"
NAVY = "123B5D"
TEAL = "0F766E"


def font(size, bold=False):
    for name in ["arial.ttf", "calibri.ttf"]:
        try:
            return ImageFont.truetype(name, size=size)
        except OSError:
            pass
    return ImageFont.load_default()


def rounded_box(draw, xy, fill, outline, label, text_fill="#172033", radius=22, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)
    x1, y1, x2, y2 = xy
    f = font(28, True)
    box = draw.textbbox((0, 0), label, font=f)
    draw.text(((x1 + x2 - box[2]) / 2, (y1 + y2 - box[3]) / 2 - 2), label, font=f, fill=text_fill)


def make_use_case(path):
    img = Image.new("RGB", (1800, 1120), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((420, 70, 1380, 1050), radius=30, fill="#F7FAFC", outline="#123B5D", width=5)
    title = "SMART WATER MONITORING CONSUMPTION SYSTEM"
    d.text((900, 100), title, anchor="ma", font=font(33, True), fill="#123B5D")

    actors = [(95, 280, "Municipal\nAdministrator"), (95, 710, "Household\nUser"), (1535, 500, "Smart Meter /\nIoT Device")]
    for x, y, label in actors:
        d.ellipse((x + 54, y, x + 104, y + 50), outline="#123B5D", width=5)
        d.line((x + 79, y + 50, x + 79, y + 145), fill="#123B5D", width=5)
        d.line((x + 25, y + 88, x + 133, y + 88), fill="#123B5D", width=5)
        d.line((x + 79, y + 145, x + 35, y + 208), fill="#123B5D", width=5)
        d.line((x + 79, y + 145, x + 123, y + 208), fill="#123B5D", width=5)
        d.multiline_text((x + 79, y + 220), label, anchor="ma", align="center", font=font(25, True), fill="#172033", spacing=3)

    admin_cases = [
        (520, 220, 860, 305, "Manage households"),
        (520, 335, 860, 420, "Manage meters & rates"),
        (520, 450, 860, 535, "Generate billing"),
        (520, 565, 860, 650, "Review alerts & reports"),
    ]
    user_cases = [
        (930, 700, 1280, 785, "View live consumption"),
        (930, 815, 1280, 900, "Review history & bills"),
        (930, 930, 1280, 1015, "Read alerts & notices"),
    ]
    device_cases = [
        (930, 250, 1280, 335, "Activate device"),
        (930, 410, 1280, 495, "Upload water reading"),
        (930, 570, 1280, 655, "Report meter health"),
    ]
    for case in admin_cases + user_cases + device_cases:
        rounded_box(d, case[:4], "#EAF3F8", "#2F6F95", case[4])

    for _, y1, _, y2, _ in admin_cases:
        d.line((230, 390, 520, (y1 + y2) // 2), fill="#718096", width=3)
    for _, y1, _, y2, _ in user_cases:
        d.line((230, 820, 930, (y1 + y2) // 2), fill="#718096", width=3)
    for _, y1, _, y2, _ in device_cases:
        d.line((1535, 600, 1280, (y1 + y2) // 2), fill="#718096", width=3)

    img.save(path, quality=95)


def remove_outside_chapter_four(doc):
    body = doc._element.body
    children = list(body)
    start = end = None
    def paragraph_text(child):
        return "".join(node.text or "" for node in child.xpath(".//w:t")).strip().upper()
    for idx, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        text = paragraph_text(child)
        if text == "CHAPTER IV":
            start = idx
        if text == "CHAPTER V" and start is not None:
            end = idx
            break
    if start is None or end is None:
        raise RuntimeError("Could not locate Chapter IV and Chapter V boundaries")
    for child in children[:start]:
        body.remove(child)
    # Refresh after deleting the leading material.
    children = list(body)
    end = next(i for i, c in enumerate(children) if c.tag == qn("w:p") and paragraph_text(c) == "CHAPTER V")
    for child in children[end:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    # Remove the page-break paragraph that originally introduced Chapter V.
    for child in reversed(list(body)):
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag == qn("w:p") and not paragraph_text(child):
            body.remove(child)
            continue
        break


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def insert_table_before(target, headers, rows, widths):
    doc = target._parent._parent
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for i, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[i]
        cell.width = Inches(width)
        cell.text = header
        set_cell_margins(cell)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), NAVY)
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(9)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row in rows:
        cells = table.add_row().cells
        for i, (value, width) in enumerate(zip(row, widths)):
            cells[i].width = Inches(width)
            cells[i].text = value
            set_cell_margins(cells[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.size = Pt(8.5)
        cant_split = OxmlElement("w:cantSplit")
        table.rows[-1]._tr.get_or_add_trPr().append(cant_split)
    target._p.addprevious(table._tbl)
    return table


def insert_figure_before(target, path, caption, alt_text, width=6.25):
    p = target.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", alt_text)
    cp = target.insert_paragraph_before(caption)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.style = "Caption"
    for r in cp.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(9)
        r.font.italic = True


def find_paragraph(doc, exact):
    return next(p for p in doc.paragraphs if p.text.strip() == exact)


def build():
    WORK.mkdir(parents=True, exist_ok=True)
    make_use_case(USE_CASE)
    doc = Document(SOURCE)
    remove_outside_chapter_four(doc)

    title = find_paragraph(doc, "METHODOLOGY, SYSTEM DESIGN, AND IMPLEMENTATION")
    title.text = "METHODOLOGY"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)

    intro = doc.paragraphs[2]
    intro.text = (
        "This chapter presents the methodology used in developing the Smart Water Monitoring Consumption System. "
        "Following the sequence of the supplied Chapter 4 reference, it discusses requirements and use cases, system "
        "design and user interfaces, the entity relationship diagram, the data dictionary, development and testing, "
        "the implementation plan, encountered challenges, and implementation results. All technical statements are "
        "based on the implemented Laravel application, ESP32 firmware, database migrations, routes, services, views, "
        "and automated tests of the Smart Water project."
    )
    intro.paragraph_format.first_line_indent = Inches(0.5)
    intro.paragraph_format.line_spacing = 1.5
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    target_42 = find_paragraph(doc, "4.2 System Design, Software, and Processes")
    h = target_42.insert_paragraph_before("4.1.4 Use-Case Diagram", style="Heading 2")
    insert_figure_before(
        target_42,
        USE_CASE,
        "Figure 4.1. Use-case overview of the Smart Water Monitoring Consumption System.",
        "Use-case overview showing the municipal administrator, household user, and smart meter actors with their primary system interactions.",
        6.25,
    )
    # Renumber the existing figures because the use-case figure is inserted first.
    for p in doc.paragraphs:
        if p.text.startswith("Figure 4.1. Smart Water"):
            p.text = p.text.replace("Figure 4.1.", "Figure 4.2.", 1)
        elif p.text.startswith("Figure 4.2. Device commissioning"):
            p.text = p.text.replace("Figure 4.2.", "Figure 4.3.", 1)
        elif p.text.startswith("Figure 4.3. Core entity"):
            p.text = p.text.replace("Figure 4.3.", "Figure 4.6.", 1)

    target_43 = find_paragraph(doc, "4.3 Entity Relationship Diagram")
    target_43.insert_paragraph_before("4.2.3 Output and User Interface Design", style="Heading 2")
    p = target_43.insert_paragraph_before(
        "The output and user interface design separates household monitoring from municipal administration. "
        "Household screens prioritize current flow, consumption history, billing estimates, leak status, notifications, "
        "and account settings. Administrator screens prioritize customers, devices, rates, statements, consumption, "
        "leak incidents, reports, notifications, analytics, and audit information. The responsive Blade layouts maintain "
        "consistent navigation while preserving role-based access boundaries."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.line_spacing = 1.5
    insert_figure_before(
        target_43,
        HOUSEHOLD,
        "Figure 4.4. Household workspace interface map.",
        "Household interface map showing dashboard, live monitoring, history, estimated billing, leak detection, notifications, and account settings.",
        6.1,
    )
    insert_figure_before(
        target_43,
        ADMIN,
        "Figure 4.5. Administrator workspace interface map.",
        "Administrator interface map showing municipal customer, device, billing, consumption, leak, report, notification, analytics, and audit modules.",
        6.1,
    )
    insert_table_before(
        target_43,
        ["User Interface", "Primary Outputs", "Design Purpose"],
        [
            ("Household dashboard", "Current flow, usage, estimated bill, meter status, leak state, and recent notices", "Present essential account information at a glance."),
            ("Live monitoring and history", "Recent readings, charts, period totals, summaries, and availability states", "Support near-real-time review and time-based consumption analysis."),
            ("Billing and alerts", "Rate breakdown, statements, leak status, thresholds, and notifications", "Make charges and exceptional conditions understandable and actionable."),
            ("Administrator workspace", "Household, meter, rate, billing, consumption, alert, report, and audit records", "Centralize municipal operations while maintaining traceability."),
        ],
        [1.45, 3.05, 2.0],
    )

    # Restore a restrained academic chapter-opening style and reference-like margins.
    section = doc.sections[0]
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    for style_name in ("Normal", "Body Text"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(11)
    for p in doc.paragraphs:
        if p.style.name == "Normal" and p.text and not p.text.startswith("Figure"):
            p.paragraph_format.line_spacing = 1.5
        for r in p.runs:
            if p.style.name == "Normal":
                r.font.name = "Arial"

    props = doc.core_properties
    props.title = "Smart Water Monitoring Consumption System - Chapter IV"
    props.subject = "Methodology, system design, development, testing, and implementation"
    props.author = "Smart Water Monitoring Consumption System Project Team"
    props.keywords = "smart water, monitoring, consumption, IoT, Laravel, ESP32, methodology"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
