from pathlib import Path
import importlib.util

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\laragon\www\SmartWater_Update")
WORK = ROOT / "artifacts" / "chapter4_ref_work"
BASE_BUILDER = WORK / "build_chapter4.py"
OUT = ROOT / "artifacts" / "Smart_Water_Monitoring_Consumption_System_Chapter_4.docx"
NAVY = "123B5D"
PALE = "EAF3F8"


def regenerate_base():
    spec = importlib.util.spec_from_file_location("chapter4_base", BASE_BUILDER)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    module.build()


def para_text(child):
    return "".join(n.text or "" for n in child.xpath(".//w:t")).strip()


def find_paragraph(doc, exact):
    return next(p for p in doc.paragraphs if p.text.strip() == exact)


def set_cell_shading(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top=85, start=100, bottom=85, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_geometry(table, widths):
    table.autofit = False
    total = int(sum(widths) * 1440)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            cell.width = Inches(width)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def format_cell(cell, size=8.5, bold=False, white=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        has_numbering = p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None
        if not has_numbering:
            p.paragraph_format.left_indent = Inches(0)
            p.paragraph_format.right_indent = Inches(0)
            p.paragraph_format.first_line_indent = Inches(0)
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = bold
            if white:
                run.font.color.rgb = RGBColor(255, 255, 255)


def fresh_numbering_id(doc):
    numbering = doc.part.numbering_part.element
    abstract_id = None
    for abstract in numbering.findall(qn("w:abstractNum")):
        styles = abstract.findall(".//" + qn("w:pStyle"))
        if any(s.get(qn("w:val")) == "ListNumber" for s in styles):
            abstract_id = abstract.get(qn("w:abstractNumId"))
            break
    if abstract_id is None:
        abstract_id = "0"
    existing = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    aid = OxmlElement("w:abstractNumId")
    aid.set(qn("w:val"), str(abstract_id))
    num.append(aid)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def move_before(target, element):
    target._p.addprevious(element)


def add_caption_before(doc, target, text):
    p = doc.add_paragraph(text, style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.italic = True
    move_before(target, p._p)


def add_heading_before(doc, target, text, level=2, page_break=False):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    if page_break:
        p.paragraph_format.page_break_before = True
    move_before(target, p._p)
    return p


def add_intro_before(doc, target, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.line_spacing = 1.5
    move_before(target, p._p)


def add_labeled_text(cell, label, value):
    p = cell.paragraphs[0]
    p.clear()
    r = p.add_run(label)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(8.5)
    r2 = p.add_run(value)
    r2.font.name = "Arial"
    r2.font.size = Pt(8.5)


USE_CASES = [
    {
        "id": "UC-01", "name": "Log In to the System", "importance": "High", "actor": "Municipal Administrator; Household User", "type": "Essential",
        "scenario": "An authorized administrator or household user wants to access the correct protected workspace.",
        "brief": "Authenticates a registered account and directs the user to the workspace permitted by the account role and status.",
        "trigger": "The user opens the login page and submits an email address and password.",
        "flow": ["The user opens the Smart Water login page.", "The user enters a registered email address and password.", "The system validates the submitted credentials.", "The system verifies that the account is active, approved, and permitted to sign in.", "The system issues the authenticated session token.", "The system redirects the administrator or household user to the corresponding workspace."],
        "sub": "The user may use the password-recovery function when the account is eligible.",
        "alt": "Invalid credentials return a validation message. Suspended, archived, inactive, incomplete, or unauthorized accounts are denied access without exposing protected data.",
    },
    {
        "id": "UC-02", "name": "Manage Household Accounts", "importance": "High", "actor": "Municipal Administrator", "type": "Essential",
        "scenario": "The administrator wants to create, locate, update, suspend, deactivate, or archive a household account.",
        "brief": "Maintains municipal household identities and their account status, address, contact information, and meter assignment readiness.",
        "trigger": "The administrator selects Customer Management from the administrative navigation.",
        "flow": ["The administrator signs in and opens Customer Management.", "The system displays searchable and filterable household records.", "The administrator selects an existing household or chooses to create a new record.", "The administrator enters or revises the permitted household details.", "The system validates required address, email, role, status, and account-number rules.", "The system saves the change and records the administrative activity."],
        "sub": "The administrator may suspend, deactivate, archive, reactivate, or unassign a household meter according to operational need.",
        "alt": "Duplicate email or account numbers, missing required fields, and unauthorized role changes are rejected with a specific validation response.",
    },
    {
        "id": "UC-03", "name": "Register and Commission a Smart Meter", "importance": "High", "actor": "Municipal Administrator; Smart Meter", "type": "Essential",
        "scenario": "The municipality wants to register a physical meter, assign it to an eligible household, and activate the field device.",
        "brief": "Creates the municipal asset record and exchanges a short-lived activation code for a scoped device credential.",
        "trigger": "The administrator selects Device Management and chooses Register Device.",
        "flow": ["The administrator enters the meter asset, serial, installation, and household details.", "The system validates uniqueness and household eligibility.", "The system stores the pending device and issues a one-time expiring activation code.", "The technician enters the code on the ESP32/M5Stack device.", "The device submits its UUID, hardware identity, and firmware version.", "The server binds the hardware identity, activates the device, and issues a scoped credential.", "The system displays the device as active and assigned."],
        "sub": "An administrator may rotate credentials, reassign the meter, schedule calibration, suspend it, or decommission the asset.",
        "alt": "Expired or reused codes, duplicate serial numbers, mismatched hardware identity, and ineligible household assignments are rejected.",
    },
    {
        "id": "UC-04", "name": "Upload a Water Reading", "importance": "High", "actor": "Smart Meter / IoT Device", "type": "Essential",
        "scenario": "An activated meter wants to submit a measured consumption interval to the server.",
        "brief": "Authenticates the device, validates interval telemetry, saves the reading once, and returns authoritative totals and alert state.",
        "trigger": "The firmware reaches the configured upload interval with a pending measurement.",
        "flow": ["The sensor counts valid flow pulses during the sampling interval.", "The firmware calculates flow rate and interval liters.", "The meter creates a stable reading identifier and stores the pending state.", "The meter submits the reading with its scoped device token.", "The server derives the owning device and household from the authenticated credential.", "The server validates and stores the reading idempotently.", "The server evaluates thresholds and returns the authoritative total and status.", "After acknowledgement, the meter clears the pending interval and advances its sequence."],
        "sub": "The device also reports Wi-Fi, server connectivity, firmware, and last-seen health information.",
        "alt": "If transmission fails, the pending reading remains stored for retry. Duplicate reading identifiers return the existing result without double-counting consumption.",
    },
    {
        "id": "UC-05", "name": "View Live Water Consumption", "importance": "High", "actor": "Household User", "type": "Essential",
        "scenario": "A household user wants to see the latest meter state and near-real-time water use.",
        "brief": "Presents the latest owned reading, current flow, interval consumption, estimated billing state, leak state, and meter availability.",
        "trigger": "The household user selects Dashboard or Live Monitoring.",
        "flow": ["The household user signs in.", "The user opens the dashboard or Live Monitoring page.", "The browser requests the latest household-scoped readings.", "The server verifies ownership and retrieves current device and consumption state.", "The interface displays current flow, liters used, estimated bill, meter connection, leak state, and recent activity.", "The interface refreshes on the configured near-real-time cadence."],
        "sub": "The user may filter recent readings by date and displayed row count.",
        "alt": "When no assigned meter or recent reading exists, the interface presents a specific unavailable state instead of reporting misleading zero consumption.",
    },
    {
        "id": "UC-06", "name": "Review Consumption History", "importance": "High", "actor": "Household User; Municipal Administrator", "type": "Essential",
        "scenario": "An authorized user wants to review water consumption for a daily, weekly, monthly, or custom period.",
        "brief": "Aggregates owned readings into period totals, summaries, comparisons, and chart-ready series while preserving missing-data semantics.",
        "trigger": "The user opens Consumption History or administrative Consumption Monitoring.",
        "flow": ["The user selects a reporting period or valid custom date range.", "The browser submits the requested range and filters.", "The server applies role and household scope.", "The consumption service retrieves eligible readings and handles interval or cumulative measurement semantics.", "The service computes period totals, summaries, and comparison values.", "The interface displays cards, a trend chart, and detailed history."],
        "sub": "The administrator may search for a household and review its scoped monitoring record.",
        "alt": "Invalid ranges are rejected. Missing readings remain unavailable rather than being inferred as zero; cumulative meter resets are handled with safe differences.",
    },
    {
        "id": "UC-07", "name": "Configure Water Rates and Alert Rules", "importance": "High", "actor": "Municipal Administrator", "type": "Essential",
        "scenario": "The municipality wants to define effective billing and monitoring rules.",
        "brief": "Maintains effective-dated water rates, fixed charges, subsidies, currency, consumption thresholds, and leak notification behavior.",
        "trigger": "The administrator opens Water Rate Management.",
        "flow": ["The administrator reviews the active municipal or household-specific configuration.", "The administrator enters the rate, fixed charge, subsidy, currency, and effective date.", "The administrator configures consumption and leak thresholds and notification behavior.", "The system validates numeric ranges, required values, and effective dates.", "The system stores the configuration and records the administrative activity.", "Subsequent estimates and statements use the latest applicable active configuration."],
        "sub": "A household-specific configuration may override the municipal default where permitted.",
        "alt": "Negative charges, invalid units, overlapping or invalid effective dates, and unauthorized changes are rejected.",
    },
    {
        "id": "UC-08", "name": "Generate a Billing Statement", "importance": "High", "actor": "Municipal Administrator", "type": "Essential",
        "scenario": "The administrator wants to generate a finalized monthly statement for an eligible household.",
        "brief": "Aggregates billable readings, applies the applicable charge rules, and preserves a historical monthly billing snapshot.",
        "trigger": "The administrator selects a household and billing month in Billing Management.",
        "flow": ["The administrator opens Billing Management and chooses Generate Statement.", "The administrator selects the household and billing month.", "The server validates the period and applicable rate configuration.", "The billing service aggregates eligible consumption and converts liters to cubic meters.", "The service applies the effective rate, fixed charge, subsidy, and opening balance.", "The system stores the statement with current charge, remaining balance, status, and charge snapshot.", "The administrator reviews the finalized billing record."],
        "sub": "Authorized staff may update payment status and record the amount paid according to the implemented billing workflow.",
        "alt": "A duplicate month, absent billable readings, or unavailable applicable rate produces a specific response and no misleading finalized statement.",
    },
    {
        "id": "UC-09", "name": "Review and Resolve a Leak Alert", "importance": "High", "actor": "Municipal Administrator", "type": "Essential",
        "scenario": "The administrator wants to investigate a possible leak and record the municipal response.",
        "brief": "Turns detected threshold conditions into a traceable incident that can be acknowledged, resolved, marked false, or reopened.",
        "trigger": "A water reading satisfies configured leak conditions or the administrator opens Leak Monitoring.",
        "flow": ["The system creates or updates the household leak incident.", "The system associates relevant readings and a household notification.", "The administrator opens Leak Monitoring and reviews the affected household, latest flow, detection times, and notice state.", "The administrator records follow-up notes and classifies the incident.", "The administrator resolves the incident or marks it as a false alarm.", "The system stores timestamps, notes, and the updated incident state."],
        "sub": "A resolved or false-alarm incident may be reopened when new evidence requires additional follow-up.",
        "alt": "Unauthorized users cannot access administrative incidents. Missing or invalid incident identifiers return a not-found or validation response.",
    },
    {
        "id": "UC-10", "name": "Manage Notifications, Reports, and Audit Records", "importance": "Medium", "actor": "Municipal Administrator", "type": "Supporting",
        "scenario": "The administrator wants to communicate official notices and review operational records.",
        "brief": "Supports notification delivery lifecycle, filtered operational reports and exports, analytics, and traceable administrative activity.",
        "trigger": "The administrator opens Notifications, Reports, Analytics, or Audit Logs.",
        "flow": ["The administrator selects the required administrative module.", "The administrator chooses an audience, report type, range, or audit filter.", "The system validates authorization and input parameters.", "The system retrieves role-scoped records or creates the official notice.", "The interface displays delivery/read state, report totals, analytic summaries, or audit context.", "The administrator may export supported reports or retry/archive eligible notification records."],
        "sub": "Households can read or remove notices from their own view without erasing the municipality's source record.",
        "alt": "Unsupported ranges, empty export selections, invalid audiences, and unauthorized audit access are rejected with clear responses.",
    },
]


# Use cases must correspond one-to-one with the labels and actors in Figure 4.1.
USE_CASES = [
    {
        "id": "UC-01", "name": "Manage Households", "importance": "High", "actor": "Municipal Administrator", "type": "Essential",
        "scenario": "The administrator wants to create, review, update, or change the operational status of a household account.",
        "brief": "Maintains household identity, service address, contact information, account number, access status, and meter-assignment readiness.",
        "trigger": "The administrator selects Customer Management from the administrative workspace.",
        "flow": ["The administrator signs in and opens Customer Management.", "The system displays searchable and filterable household records.", "The administrator creates a new household or selects an existing record.", "The administrator enters or updates the permitted household information.", "The system validates required fields, unique account details, role, and status.", "The system saves the household record and logs the administrative activity."],
        "sub": "The administrator may suspend, deactivate, archive, reactivate, or unassign the household meter when authorized.",
        "alt": "Duplicate emails or account numbers, missing required fields, and unauthorized role changes are rejected with a validation response.",
    },
    {
        "id": "UC-02", "name": "Manage Meters and Rates", "importance": "High", "actor": "Municipal Administrator", "type": "Essential",
        "scenario": "The administrator wants to manage municipal smart meters and configure the water rates and monitoring rules applied to households.",
        "brief": "Combines meter registration, assignment, maintenance, credential control, and effective-dated water-rate and alert configuration.",
        "trigger": "The administrator opens Device Management or Water Rate Management.",
        "flow": ["The administrator opens the required meter or rate module.", "The system displays the current devices or active billing configurations.", "For a meter, the administrator registers, assigns, reassigns, suspends, rotates credentials, or decommissions the asset.", "For a rate, the administrator enters the water rate, fixed charge, subsidy, currency, effective date, and alert thresholds.", "The system validates identifiers, household eligibility, numeric ranges, and effective dates.", "The system saves the approved meter or rate configuration and records the activity."],
        "sub": "The administrator may review meter health, calibration dates, activation state, last-seen status, or household-specific rate overrides.",
        "alt": "Duplicate assets, ineligible household assignments, invalid rates, negative charges, and unauthorized changes are rejected.",
    },
    {
        "id": "UC-03", "name": "Generate Billing", "importance": "High", "actor": "Municipal Administrator", "type": "Essential",
        "scenario": "The administrator wants to generate a monthly billing statement for an eligible household.",
        "brief": "Aggregates billable consumption, applies the applicable rate configuration, and preserves a historical charge snapshot.",
        "trigger": "The administrator chooses Generate Statement in Billing Management.",
        "flow": ["The administrator opens Billing Management.", "The administrator selects the household and billing month.", "The system validates the period, household, billable readings, and applicable rate.", "The billing service totals liters and converts the result to cubic meters.", "The service applies the water rate, fixed charge, subsidy, and opening balance.", "The system stores the statement, charge breakdown, remaining balance, and payment status.", "The administrator reviews the generated billing record."],
        "sub": "Authorized staff may record payments and update the statement status through the implemented billing workflow.",
        "alt": "Duplicate billing months, missing billable readings, or unavailable rate configurations produce a specific response without creating a misleading statement.",
    },
    {
        "id": "UC-04", "name": "Review Alerts and Reports", "importance": "High", "actor": "Municipal Administrator", "type": "Essential",
        "scenario": "The administrator wants to review leak or consumption alerts and examine municipal operational reports.",
        "brief": "Supports alert investigation, incident follow-up, analytics, filtered reports, and supported data exports.",
        "trigger": "The administrator opens Leak Monitoring, Reports, Analytics, or the administrative notification view.",
        "flow": ["The administrator selects the alert or report module.", "The administrator chooses the affected household, report type, date range, or filter.", "The system validates authorization and input parameters.", "The system retrieves the related readings, alerts, billing information, and operational summaries.", "The administrator reviews the evidence and records follow-up notes when an incident is involved.", "The administrator resolves, marks false, or reopens an alert, or exports a supported report."],
        "sub": "The administrator may review notification delivery state, analytic summaries, and audit context related to an operational event.",
        "alt": "Invalid ranges, unsupported exports, missing incident identifiers, and unauthorized access return clear validation or access responses.",
    },
    {
        "id": "UC-05", "name": "View Live Consumption", "importance": "High", "actor": "Household User", "type": "Essential",
        "scenario": "The household user wants to view the latest water consumption and meter state.",
        "brief": "Displays the latest owned reading, current flow, interval consumption, estimated billing state, leak state, and meter availability.",
        "trigger": "The household user opens the Dashboard or Live Monitoring page.",
        "flow": ["The household user signs in.", "The user opens Dashboard or Live Monitoring.", "The browser requests the latest household-scoped readings.", "The server verifies ownership and retrieves the current meter and consumption state.", "The interface displays current flow, liters used, estimated bill, leak status, meter connection, and recent activity.", "The interface refreshes according to the configured near-real-time interval."],
        "sub": "The user may filter recent readings by date or displayed row count.",
        "alt": "When no active meter or recent reading exists, the interface displays a specific unavailable state rather than misleading zero consumption.",
    },
    {
        "id": "UC-06", "name": "Review History and Bills", "importance": "High", "actor": "Household User", "type": "Essential",
        "scenario": "The household user wants to review past consumption and current or finalized billing information.",
        "brief": "Combines daily, weekly, monthly, or custom consumption history with estimated billing details and monthly statements.",
        "trigger": "The user opens Consumption History or Estimated Billing.",
        "flow": ["The user selects the consumption-history or billing page.", "For history, the user selects a period or valid custom date range.", "The server retrieves only the household's readings and computes totals, comparisons, and chart data.", "For billing, the server selects the applicable rate and eligible consumption information.", "The interface displays usage summaries, trends, the estimated charge breakdown, and finalized statement history.", "The user reviews the selected history or billing record."],
        "sub": "The household may switch among daily, weekly, and monthly views and inspect individual monthly statements.",
        "alt": "Invalid ranges are rejected. Missing readings or rate information remain unavailable rather than being interpreted as zero usage or zero charges.",
    },
    {
        "id": "UC-07", "name": "Read Alerts and Notices", "importance": "Medium", "actor": "Household User", "type": "Essential",
        "scenario": "The household user wants to read leak, consumption, billing, meter, or account notifications.",
        "brief": "Presents household-scoped alerts and official notices while preserving the municipality's source and delivery records.",
        "trigger": "The user opens Notifications or selects a recent alert from the household workspace.",
        "flow": ["The household user signs in and opens Notifications.", "The system retrieves notices owned by the household account.", "The interface displays each notice's type, title, message, date, and read state.", "The user opens a selected alert or notice.", "The system marks the notice as read when applicable.", "The user may remove the notice from the household view."],
        "sub": "Leak-related notices may direct the user to the Leak Detection page for the latest meter signal and response status.",
        "alt": "A household cannot access another account's notices. User-side removal hides the notice without deleting the municipality's source record.",
    },
    {
        "id": "UC-08", "name": "Activate Device", "importance": "High", "actor": "Smart Meter / IoT Device", "type": "Essential",
        "scenario": "A registered M5Stack smart meter wants to exchange its one-time code for a scoped device credential.",
        "brief": "Binds the physical hardware to the municipal device record and activates authenticated telemetry access.",
        "trigger": "The technician enters the administrator-issued activation code on the registered device.",
        "flow": ["The administrator has already registered the meter and issued a short-lived activation code.", "The device submits its UUID, activation code, hardware identity, and firmware version.", "The server verifies that the code is valid, unused, and unexpired.", "The server checks device status and hardware-binding rules.", "The server binds the hardware identity and marks the device active.", "The server issues a scoped device credential.", "The M5Stack stores the credential and becomes ready to upload telemetry."],
        "sub": "An administrator may later rotate the credential or reissue activation after an authorized lifecycle change.",
        "alt": "Expired, reused, incorrect, or mismatched activation information is rejected and no device credential is issued.",
    },
    {
        "id": "UC-09", "name": "Upload Water Reading", "importance": "High", "actor": "Smart Meter / IoT Device", "type": "Essential",
        "scenario": "An activated smart meter wants to submit a measured water-consumption interval.",
        "brief": "Authenticates the device, validates telemetry, stores the reading once, and returns authoritative totals and alert state.",
        "trigger": "The firmware reaches the configured upload interval with a pending measurement.",
        "flow": ["The sensor counts valid flow pulses during the sampling interval.", "The firmware calculates flow rate and interval liters.", "The device creates a stable reading identifier and stores the pending state.", "The meter submits the reading with its scoped credential.", "The server derives the device and household from the authenticated identity.", "The server validates and stores the reading idempotently.", "The server evaluates thresholds and returns the authoritative total and leak state.", "After acknowledgement, the meter clears the pending interval and advances its sequence."],
        "sub": "The payload may include device name, monitoring state, Wi-Fi strength, and connectivity information.",
        "alt": "When transmission fails, the pending reading remains available for retry. Duplicate reading identifiers do not double-count consumption.",
    },
    {
        "id": "UC-10", "name": "Report Meter Health", "importance": "Medium", "actor": "Smart Meter / IoT Device", "type": "Supporting",
        "scenario": "The smart meter wants to report its connection, firmware, and operational health to the server.",
        "brief": "Updates municipal visibility of device connectivity, last-seen activity, telemetry status, firmware version, and relevant diagnostics.",
        "trigger": "The device activates, uploads telemetry, reconnects to the network, or reaches a health-reporting interval.",
        "flow": ["The M5Stack evaluates its Wi-Fi and server connection state.", "The device prepares its hardware identity, firmware version, monitoring state, signal strength, and health metadata.", "The device submits the health information using its scoped credential.", "The server validates the authenticated device identity.", "The server updates last-seen, last-telemetry, firmware, and operational health fields.", "The administrator workspace displays the latest device status."],
        "sub": "Health information may accompany a water-reading upload to reduce unnecessary network requests.",
        "alt": "Invalid credentials or mismatched hardware identity are rejected. During network loss, the device retains its state and reports again after reconnection.",
    },
]


DATA_DICTIONARY = {
    "users": [
        ("id", "BIGINT", "20", "Primary key of the user account."), ("name", "VARCHAR", "255", "Display name of the administrator, household, or device account."),
        ("email", "VARCHAR", "255", "Unique login or account email address."), ("role", "VARCHAR", "255", "Authorization role: administrator, household user, or device."),
        ("account_number", "VARCHAR", "255", "Standardized municipal account identifier."), ("address", "VARCHAR", "255", "Required service address for household accounts."),
        ("phone", "VARCHAR", "255", "Optional contact number."), ("is_active", "BOOLEAN", "1", "Indicates whether the account can remain operational."),
        ("email_verified_at", "DATETIME", "-", "Date and time the email was verified."), ("password", "VARCHAR", "255", "Hashed interactive-account password."),
        ("remember_token", "VARCHAR", "100", "Framework token used for remembered sessions."), ("created_by_admin", "BOOLEAN", "1", "Indicates municipal administrator provisioning."),
        ("assigned_user_id", "BIGINT", "20", "Household linked to a legacy or supporting device account."), ("created_by", "BIGINT", "20", "Administrator who created the account."),
        ("approved_at", "DATETIME", "-", "Date and time the household account was approved."), ("account_status", "VARCHAR", "255", "Lifecycle state such as active, suspended, archived, or deactivated."),
        ("created_at", "DATETIME", "-", "Record creation timestamp."), ("updated_at", "DATETIME", "-", "Record update timestamp."),
    ],
    "devices": [
        ("id", "BIGINT", "20", "Primary key of the municipal device record."), ("uuid", "VARCHAR", "255", "Unique server-issued device identifier."),
        ("asset_number", "VARCHAR", "255", "Municipal equipment inventory number."), ("serial_number", "VARCHAR", "255", "Unique manufacturer or field serial number."),
        ("name", "VARCHAR", "255", "Human-readable meter name."), ("household_id", "BIGINT", "20", "Currently assigned household; nullable while in inventory."),
        ("auth_user_id", "BIGINT", "20", "Limited authentication identity associated with the device."), ("hardware_id", "VARCHAR", "255", "Bound ESP32 hardware identity."),
        ("firmware_version", "VARCHAR", "255", "Firmware version reported during activation or telemetry."), ("status", "VARCHAR", "255", "Lifecycle state such as pending, active, suspended, or decommissioned."),
        ("maintenance_status", "VARCHAR", "255", "Operational or maintenance condition."), ("installation_address", "VARCHAR", "255", "Physical service location of the meter."),
        ("installed_at", "DATETIME", "-", "Installation date and time."), ("last_seen_at", "DATETIME", "-", "Latest successful device contact."),
        ("last_telemetry_at", "DATETIME", "-", "Latest accepted telemetry timestamp."), ("calibrated_at", "DATETIME", "-", "Most recent calibration date."),
        ("calibration_due_at", "DATETIME", "-", "Scheduled calibration due date."), ("token_rotated_at", "DATETIME", "-", "Latest credential rotation timestamp."),
        ("revoked_at", "DATETIME", "-", "Credential or device revocation timestamp."), ("activation_code_hash", "VARCHAR", "255", "Hash of the one-time activation code."),
        ("activation_expires_at", "DATETIME", "-", "Expiration of the activation code."), ("activated_at", "DATETIME", "-", "Successful activation timestamp."),
        ("notes", "TEXT", "-", "Administrative installation or maintenance notes."), ("created_at", "DATETIME", "-", "Record creation timestamp."),
        ("updated_at", "DATETIME", "-", "Record update timestamp."),
    ],
    "water_readings": [
        ("id", "BIGINT", "20", "Primary key of the telemetry record."), ("user_id", "BIGINT", "20", "Household owner of the reading."),
        ("device_id", "BIGINT", "20", "Authenticated meter that submitted the reading."), ("reading_id", "VARCHAR", "255", "Stable idempotency identifier generated by the device."),
        ("flow_rate", "DECIMAL", "18,6", "Measured flow rate in liters per minute."), ("liters_used", "DECIMAL", "18,6", "Water volume represented by the interval."),
        ("cubic_meter", "DECIMAL", "18,6", "Consumption converted to cubic meters."), ("estimated_bill", "DECIMAL", "12,2", "Estimated charge stored with the reading."),
        ("rate_per_cubic_meter", "DECIMAL", "12,4", "Applicable water rate used for the estimate."), ("subsidy_deduction", "DECIMAL", "12,2", "Subsidy applied to the estimate."),
        ("threshold_liters", "DECIMAL", "18,6", "Consumption threshold applicable at processing time."), ("is_threshold_exceeded", "BOOLEAN", "1", "Indicates consumption threshold exceedance."),
        ("is_possible_leak", "BOOLEAN", "1", "Indicates possible leakage under server rules."), ("leak_detection_status", "VARCHAR", "255", "Detailed leak evaluation state."),
        ("water_limit_liters", "DECIMAL", "18,6", "Configured household water limit snapshot."), ("monitoring_status", "VARCHAR", "255", "Device monitoring state reported with telemetry."),
        ("wifi_signal_strength", "INTEGER", "11", "Reported Wi-Fi signal strength."), ("device_connection_status", "VARCHAR", "255", "Meter-side connection state."),
        ("server_connection_status", "VARCHAR", "255", "Latest server acknowledgement state."), ("device_name", "VARCHAR", "255", "Reported or display device name."),
        ("device_payload", "JSON/TEXT", "-", "Additional validated device diagnostics."), ("recorded_at", "DATETIME", "-", "Measurement timestamp supplied or normalized by the server."),
        ("created_at", "DATETIME", "-", "Record creation timestamp."), ("updated_at", "DATETIME", "-", "Record update timestamp."),
    ],
    "billings": [
        ("id", "BIGINT", "20", "Primary key of the billing configuration."), ("user_id", "BIGINT", "20", "Household override owner; null indicates municipal default."),
        ("water_rate", "DECIMAL", "12,4", "Charge per cubic meter."), ("fixed_monthly_charge", "DECIMAL", "12,2", "Fixed recurring monthly charge."),
        ("subsidy_deduction", "DECIMAL", "12,2", "Configured subsidy deduction."), ("currency", "VARCHAR", "3", "Three-letter billing currency, normally PHP."),
        ("effective_date", "DATE", "-", "Date on which the configuration becomes applicable."), ("threshold_liters", "DECIMAL", "18,6", "Consumption alert threshold."),
        ("consumption_alert_enabled", "BOOLEAN", "1", "Enables consumption threshold notifications."), ("consumption_alert_unit", "VARCHAR", "255", "Unit used by the consumption threshold."),
        ("consumption_alert_behavior", "VARCHAR", "255", "Audience or handling rule for consumption notices."), ("leak_alert_enabled", "BOOLEAN", "1", "Enables leak alert evaluation."),
        ("leak_flow_rate_threshold", "DECIMAL", "18,6", "Flow-rate threshold used for possible leak detection."), ("leak_continuous_duration_minutes", "INTEGER", "11", "Required continuous duration before leak handling."),
        ("leak_alert_behavior", "VARCHAR", "255", "Audience or handling rule for leak notices."), ("quiet_hours_start", "INTEGER", "2", "Starting hour of the configured quiet period."),
        ("quiet_hours_end", "INTEGER", "2", "Ending hour of the configured quiet period."), ("is_active", "BOOLEAN", "1", "Indicates whether the configuration is applicable."),
        ("created_at", "DATETIME", "-", "Record creation timestamp."), ("updated_at", "DATETIME", "-", "Record update timestamp."),
    ],
    "billing_statements": [
        ("id", "BIGINT", "20", "Primary key of the monthly statement."), ("user_id", "BIGINT", "20", "Household billed by the statement."),
        ("billing_month", "DATE", "-", "Month represented by the statement."), ("opening_balance", "DECIMAL", "12,2", "Balance carried into the billing period."),
        ("total_liters", "DECIMAL", "18,6", "Billable water consumption in liters."), ("total_cubic_meter", "DECIMAL", "18,6", "Billable water consumption in cubic meters."),
        ("water_rate", "DECIMAL", "12,4", "Rate snapshot applied at generation."), ("fixed_monthly_charge", "DECIMAL", "12,2", "Fixed-charge snapshot applied at generation."),
        ("subsidy_deduction", "DECIMAL", "12,2", "Subsidy snapshot applied at generation."), ("currency", "VARCHAR", "3", "Statement currency."),
        ("current_charge", "DECIMAL", "12,2", "Charge calculated for the current period."), ("amount_paid", "DECIMAL", "12,2", "Payment recorded against the statement."),
        ("remaining_balance", "DECIMAL", "12,2", "Outstanding balance after payment."), ("status", "VARCHAR", "255", "Payment state such as unpaid, partial, or paid."),
        ("generated_at", "DATETIME", "-", "Date and time the statement was finalized."), ("paid_at", "DATETIME", "-", "Date and time full payment was recorded."),
        ("created_at", "DATETIME", "-", "Record creation timestamp."), ("updated_at", "DATETIME", "-", "Record update timestamp."),
    ],
    "notifications": [
        ("id", "BIGINT", "20", "Primary key of the notification."), ("user_id", "BIGINT", "20", "Household recipient; nullable for administrative or batch scope."),
        ("water_reading_id", "BIGINT", "20", "Related water reading when the notice is telemetry-driven."), ("type", "VARCHAR", "255", "Notification category."),
        ("title", "VARCHAR", "255", "Short notification heading."), ("message", "TEXT", "-", "Full notification content."),
        ("metadata", "JSON/TEXT", "-", "Structured context used by the interface or workflow."), ("is_read", "BOOLEAN", "1", "Indicates whether the recipient opened the notice."),
        ("batch_id", "VARCHAR", "255", "Identifier shared by a batch announcement."), ("delivery_status", "VARCHAR", "255", "Delivery lifecycle state."),
        ("delivery_error", "TEXT", "-", "Latest delivery failure message when applicable."), ("scheduled_at", "DATETIME", "-", "Scheduled delivery timestamp."),
        ("sent_at", "DATETIME", "-", "Actual delivery or creation timestamp."), ("archived_at", "DATETIME", "-", "Administrative archive timestamp."),
        ("user_deleted_at", "DATETIME", "-", "Household-side removal timestamp without source deletion."), ("created_at", "DATETIME", "-", "Record creation timestamp."),
        ("updated_at", "DATETIME", "-", "Record update timestamp."),
    ],
    "leak_alerts": [
        ("id", "BIGINT", "20", "Primary key of the leak incident."), ("user_id", "BIGINT", "20", "Affected household."),
        ("first_water_reading_id", "BIGINT", "20", "First reading associated with the incident."), ("latest_water_reading_id", "BIGINT", "20", "Most recent supporting reading."),
        ("notification_id", "BIGINT", "20", "Notification associated with the incident."), ("status", "VARCHAR", "255", "Incident state such as detected, acknowledged, resolved, or false alarm."),
        ("notes", "TEXT", "-", "Administrative investigation and resolution notes."), ("detected_at", "DATETIME", "-", "Initial detection timestamp."),
        ("last_detected_at", "DATETIME", "-", "Most recent detection timestamp."), ("notified_at", "DATETIME", "-", "Household notification timestamp."),
        ("acknowledged_at", "DATETIME", "-", "Administrative acknowledgement timestamp."), ("resolved_at", "DATETIME", "-", "Resolution timestamp."),
        ("created_at", "DATETIME", "-", "Record creation timestamp."), ("updated_at", "DATETIME", "-", "Record update timestamp."),
    ],
    "audit_logs": [
        ("id", "BIGINT", "20", "Primary key of the audit record."), ("user_id", "BIGINT", "20", "Authenticated actor when available."),
        ("action", "VARCHAR", "255", "Machine-readable administrative or security action."), ("auditable_type", "VARCHAR", "255", "Type of affected domain record."),
        ("auditable_id", "BIGINT", "20", "Identifier of the affected domain record."), ("description", "VARCHAR", "255", "Human-readable activity description."),
        ("metadata", "JSON/TEXT", "-", "Redacted structured activity context."), ("outcome", "VARCHAR", "255", "Result such as success or failure."),
        ("ip_address", "VARCHAR", "45", "Request IP address when retained."), ("user_agent", "TEXT", "-", "Request client or browser description."),
        ("request_id", "VARCHAR", "255", "Request correlation identifier."), ("created_at", "DATETIME", "-", "Record creation timestamp."),
        ("updated_at", "DATETIME", "-", "Record update timestamp."),
    ],
}


def create_use_case_table(doc, uc):
    table = doc.add_table(rows=8, cols=3)
    table.style = "Table Grid"
    table_geometry(table, [2.15, 1.55, 2.30])
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    top = table.rows[0].cells
    add_labeled_text(top[0], "Use Case Name: ", uc["name"])
    add_labeled_text(top[1], "ID: ", uc["id"])
    add_labeled_text(top[2], "Importance Level: ", uc["importance"])
    second = table.rows[1].cells
    add_labeled_text(second[0], "Primary Actor: ", uc["actor"])
    second[1].merge(second[2])
    add_labeled_text(second[1], "Use Case Type: ", uc["type"])
    labels = [("Scenario: ", "scenario"), ("Brief Description: ", "brief"), ("Trigger: ", "trigger")]
    for row_idx, (label, key) in enumerate(labels, start=2):
        cell = table.rows[row_idx].cells[0].merge(table.rows[row_idx].cells[2])
        add_labeled_text(cell, label, uc[key])
    normal = table.rows[5].cells[0].merge(table.rows[5].cells[2])
    normal.text = ""
    p = normal.paragraphs[0]
    r = p.add_run("Normal Flow of Events:")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(8.5)
    num_id = fresh_numbering_id(doc)
    for step in uc["flow"]:
        lp = normal.add_paragraph(step)
        apply_numbering(lp, num_id)
        lp.paragraph_format.left_indent = Inches(0.25)
        lp.paragraph_format.first_line_indent = Inches(-0.18)
        lp.paragraph_format.space_after = Pt(0)
        for run in lp.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
    sub = table.rows[6].cells[0].merge(table.rows[6].cells[2])
    add_labeled_text(sub, "Sub Flows: ", uc["sub"])
    alt = table.rows[7].cells[0].merge(table.rows[7].cells[2])
    add_labeled_text(alt, "Alternative/Exceptional Flows: ", uc["alt"])
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            format_cell(cell)
    return table


def create_dictionary_table(doc, fields):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table_geometry(table, [1.60, 1.15, 0.75, 2.50])
    for idx, title in enumerate(("Field Name", "Data Type", "Length", "Description")):
        table.rows[0].cells[idx].text = title
        set_cell_shading(table.rows[0].cells[idx], NAVY)
        format_cell(table.rows[0].cells[idx], 8.0, True, True)
        table.rows[0].cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_margins(table.rows[0].cells[idx], 55, 75, 55, 75)
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for field in fields:
        cells = table.add_row().cells
        for idx, value in enumerate(field):
            cells[idx].text = value
            format_cell(cells[idx], 7.7, idx == 0)
            set_cell_margins(cells[idx], 45, 75, 45, 75)
            if idx < 3:
                cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    return table


def delete_summary_between(doc, start_heading, end_heading, keep_first_body_paragraph=False):
    body = doc._element.body
    children = list(body)
    start = next(i for i, c in enumerate(children) if c.tag == qn("w:p") and para_text(c) == start_heading)
    end = next(i for i, c in enumerate(children) if i > start and c.tag == qn("w:p") and para_text(c) == end_heading)
    kept_body = False
    for child in children[start + 1:end]:
        if keep_first_body_paragraph and child.tag == qn("w:p") and para_text(child) and not kept_body:
            kept_body = True
            continue
        body.remove(child)


def build():
    regenerate_base()
    doc = Document(OUT)

    # Match the reference sequence: diagram first, then detailed use-case descriptions.
    delete_summary_between(doc, "4.1.3 Use-Case Description", "4.1.4 Use-Case Diagram")
    h_diagram = find_paragraph(doc, "4.1.3 Use-Case Description")
    h_diagram.text = "4.1.3 Use-Case Diagram"
    old_h = find_paragraph(doc, "4.1.4 Use-Case Diagram")
    old_h._element.getparent().remove(old_h._element)
    target_42 = find_paragraph(doc, "4.2 System Design, Software, and Processes")
    add_heading_before(doc, target_42, "4.1.4 Use-Case Descriptions", 2, True)
    add_intro_before(doc, target_42, "The following use-case descriptions adapt the detailed documentation format shown in the reference to the implemented Smart Water actors and workflows. Each specification identifies the responsible actor, scenario, trigger, normal sequence, supporting path, and exceptional conditions.")
    for idx, uc in enumerate(USE_CASES, start=1):
        h = add_heading_before(doc, target_42, f"4.1.4.{idx} {uc['name']}", 3, idx > 1)
        table = create_use_case_table(doc, uc)
        move_before(target_42, table._tbl)
        add_caption_before(doc, target_42, f"Table 4.1.{idx} {uc['name']}")

    # Replace the compact dictionary summary with field-level table definitions.
    delete_summary_between(doc, "4.4 Data Dictionary", "4.5 Development and Testing", keep_first_body_paragraph=True)
    target_45 = find_paragraph(doc, "4.5 Development and Testing")
    display_names = {
        "users": "Users", "devices": "Devices", "water_readings": "Water Readings", "billings": "Billing Configurations",
        "billing_statements": "Billing Statements", "notifications": "Notifications", "leak_alerts": "Leak Alerts", "audit_logs": "Audit Logs",
    }
    for idx, (table_name, fields) in enumerate(DATA_DICTIONARY.items(), start=1):
        add_heading_before(doc, target_45, f"4.4.{idx} {display_names[table_name]} Table", 2, True)
        intro = f"The {table_name} table stores {display_names[table_name].lower()} data used by the implemented Smart Water application."
        add_intro_before(doc, target_45, intro)
        table = create_dictionary_table(doc, fields)
        move_before(target_45, table._tbl)
        add_caption_before(doc, target_45, f"Table 4.4.{idx} Database Table of {display_names[table_name]}")

    props = doc.core_properties
    props.title = "Smart Water Monitoring Consumption System - Detailed Chapter IV"
    props.subject = "Detailed methodology, use cases, interface design, database dictionary, testing, and implementation"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
