from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(r"C:\laragon\capstone\SmartWater_Update")
REFERENCE = ROOT / "artifacts" / "Smart_Water_Monitoring_Consumption_System_Chapters_4_and_5.docx"
OUTPUT = ROOT / "artifacts" / "Smart_Water_Monitoring_Consumption_System_Chapter_6.docx"


def clear_body_keep_sectpr(document):
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def font(run, size=None, bold=None, italic=None, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_text(document, text, style="Normal", align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             before=0, after=0, first_indent=None, size=12, bold=None, italic=None):
    paragraph = document.add_paragraph(style=style)
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    if first_indent is not None:
        paragraph.paragraph_format.first_line_indent = Pt(first_indent)
    run = paragraph.add_run(text)
    font(run, size=size, bold=bold, italic=italic)
    return paragraph


def add_numbered_recommendation(document, title, description):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(4)
    lead = paragraph.add_run(title + ". ")
    font(lead, size=12, bold=True)
    body = paragraph.add_run(description)
    font(body, size=12)
    return paragraph


doc = Document(REFERENCE)
clear_body_keep_sectpr(doc)

# Cover page: retained typography and whitespace pattern.
add_text(doc, "SMART WATER MONITORING\nCONSUMPTION SYSTEM", align=WD_ALIGN_PARAGRAPH.CENTER,
         before=115, size=20, bold=True)
add_text(doc, "CHAPTER VI", align=WD_ALIGN_PARAGRAPH.CENTER, before=25, size=14, bold=True)
add_text(doc, "Summary, Conclusions, and Recommendations\nCalape, Bohol",
         align=WD_ALIGN_PARAGRAPH.CENTER, before=120, size=12, italic=True)
add_text(doc, "Prepared from the implemented Laravel and ESP32 project\nOctober 2026",
         align=WD_ALIGN_PARAGRAPH.CENTER, before=100, size=12)
cover_end = doc.add_paragraph()
cover_end.add_run().add_break(WD_BREAK.PAGE)

# Chapter opener.
add_text(doc, "CHAPTER VI", align=WD_ALIGN_PARAGRAPH.CENTER, before=24, after=4, size=14, bold=True)
add_text(doc, "SUMMARY, CONCLUSIONS, AND RECOMMENDATIONS", align=WD_ALIGN_PARAGRAPH.CENTER,
         after=20, size=14, bold=True)
add_text(
    doc,
    "This chapter synthesizes the outcomes of the Smart Water Monitoring Consumption System and presents conclusions based on the implemented application, device workflow, and available verification evidence. It also identifies the actions required before the system can be treated as ready for a wider municipal deployment.",
)

add_text(doc, "6.1 Summary", style="Heading 1", align=None, size=12)
add_text(
    doc,
    "The project developed a web-based Smart Water Monitoring Consumption System that connects an ESP32-S3/M5Stack CoreS3 flow-meter device with a Laravel application for households and municipal personnel. The device measures water flow, retains pending interval readings during temporary connection failures, and submits readings using a stable identifier. The server validates the device identity and household association before storing readings and producing consumption, alert, and billing information.",
)
add_text(
    doc,
    "For household users, the system provides access to near-real-time monitoring, consumption history, estimated billing, finalized statements and receipts, leak-related notices, account maintenance, language labels, low-bandwidth settings, and service requests. For administrators, it provides controlled household onboarding, device registration and commissioning, water-rate configuration, billing statement management, leak-incident handling, notifications, reports, analytics, payment records, and audit logs. These functions create a traceable path from measured water use to a household-facing record and an administrative response.",
)
add_text(
    doc,
    "The project also incorporates safeguards intended to improve data integrity and security. These include role-scoped access, one-time device activation codes, device credential revocation, stable reading identifiers to reduce duplicate counting, server-authoritative ownership, retained billing snapshots, and audit-log redaction for sensitive values. Automated feature tests and a production frontend build provide evidence that major software behaviors have been exercised; however, software verification does not replace physical-meter calibration, field validation, or formal operational acceptance.",
)

add_text(doc, "6.2 Conclusions", style="Heading 1", align=None, size=12)
add_text(
    doc,
    "Based on the implemented features and the documented verification activities, the Smart Water Monitoring Consumption System achieves its core functional objective of integrating water-flow telemetry with household monitoring and municipal management. It provides a structured means to collect meter readings, associate them with the correct device and household, present current and historical consumption, estimate charges, and support operational follow-up for unusual consumption and possible leaks.",
)
add_text(
    doc,
    "The system demonstrates that smart-meter data can be made more useful when it is paired with role-based interfaces and explicit operational workflows. Households can review information relevant to their own accounts, while administrators can manage devices, rates, statements, incidents, and reports without relying on the household interface for municipal actions. This separation supports accountability and reduces the risk of exposing records or controls outside their intended roles.",
)
add_text(
    doc,
    "The system's reliability measures strengthen the design but do not establish field accuracy on their own. Pending-reading storage, acknowledgement-based clearing, and idempotent reading identifiers address common connectivity and duplicate-submission risks. Nevertheless, the actual measurement accuracy of each installed meter depends on calibration, installation conditions, sensor characteristics, pressure, and network behavior. The system should therefore be regarded as suitable for controlled pilot use and continued refinement, not as fully production-ready solely on the basis of the current software evidence.",
)
add_text(
    doc,
    "Finally, billing and alert outputs are useful decision-support records only when their configuration and source readings are valid. Municipal staff must review effective rates, fixed charges, subsidies, thresholds, rounding rules, and official payment procedures before relying on system-generated estimates or statements for formal collection. Full acceptance also requires a completed, green regression run in each supported database environment and documented recovery, backup, and security checks.",
)

add_text(doc, "6.3 Recommendations", style="Heading 1", align=None, size=12)
add_numbered_recommendation(
    doc,
    "Complete regression verification",
    "Investigate and resolve any failing or environment-specific tests, especially database-dependent history and aggregation behavior. Run the complete automated suite and production build in each supported database configuration before formal acceptance.",
)
add_numbered_recommendation(
    doc,
    "Conduct calibrated field trials",
    "Compare each installed meter against a measured reference volume at low, normal, and high flow. Record the resulting calibration factor, error range, installation conditions, and acceptance threshold for every device.",
)
add_numbered_recommendation(
    doc,
    "Run a controlled pilot",
    "Deploy to a limited set of households first and measure telemetry delay, outage recovery, data completeness, billing variance, leak-alert sensitivity, false alarms, and user task completion before expanding coverage.",
)
add_numbered_recommendation(
    doc,
    "Formalize municipal billing rules",
    "Confirm the authorized rate schedule, fixed charges, subsidies, arrears, due dates, rounding, adjustments, payment posting, and dispute procedures with the responsible municipal office. Preserve approved billing configurations and statement snapshots for auditability.",
)
add_numbered_recommendation(
    doc,
    "Strengthen production operations",
    "Use HTTPS, protected environment secrets, least-privilege database accounts, scheduled backups, restoration drills, queue and scheduler monitoring, software patching, and documented incident procedures before live deployment.",
)
add_numbered_recommendation(
    doc,
    "Prepare role-based training",
    "Provide concise guides for administrators, field technicians, and household users covering meter commissioning, account security, meter-health review, alert response, statement handling, and service-request follow-up.",
)
add_numbered_recommendation(
    doc,
    "Evaluate notification expansion carefully",
    "If email or SMS is added after pilot validation, obtain consent where required and implement delivery retries, rate limits, delivery logs, opt-out controls, and privacy safeguards. Critical operational notices should remain governed by clearly defined delivery rules.",
)
add_numbered_recommendation(
    doc,
    "Use accumulated data responsibly",
    "After sufficient validated data has been collected, analyze seasonal demand, recurring leaks, meter anomalies, and conservation outcomes. Missing or delayed telemetry must remain distinguishable from zero water use in all future analyses.",
)

add_text(doc, "6.4 Chapter Summary", style="Heading 1", align=None, size=12)
add_text(
    doc,
    "Chapter VI summarized the implemented Smart Water Monitoring Consumption System, concluded that it fulfills its core software objectives, and identified the conditions required for responsible deployment. The next phase should focus on calibration, a controlled field pilot, full regression verification, formal billing approval, operational safeguards, and user training. These steps will determine whether the system can progress from a validated capstone implementation to a reliable municipal service.",
)

doc.core_properties.title = "Smart Water Monitoring Consumption System Chapter VI"
doc.core_properties.subject = "Summary Conclusions and Recommendations"
doc.core_properties.author = "Smart Water Monitoring Consumption System Project Team"
doc.save(OUTPUT)
print(OUTPUT)
