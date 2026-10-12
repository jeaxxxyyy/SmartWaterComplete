from pathlib import Path

from docx import Document
from docx.shared import Pt


source = Path(r"C:\laragon\capstone\SmartWater_Update\artifacts\Smart_Water_Monitoring_Consumption_System_Chapters_4_and_5.docx")
output = Path(r"C:\laragon\capstone\SmartWater_Update\Smart-Water-Consumption-Monitoring-Updated-Chapters-IV-V.docx")
document = Document(source)

replacements = {
    "Smart Water Monitoring Consumption System": "Smart Water Consumption Monitoring System",
    "SMART WATER MONITORING CONSUMPTION SYSTEM": "SMART WATER CONSUMPTION MONITORING SYSTEM",
    "ZJ-S201": "YF-S201",
    "M5Stack device": "M5Stack CoreS3 device",
    "137 passing automated tests demonstrate broad functional coverage, while the remaining failing test identifies a specific acceptance item rather than an undocumented weakness.": (
        "181 automated tests with 1,178 assertions passed in the current verification run, demonstrating broad functional coverage of the implemented software."
    ),
    "The current codebase is close to automated-test acceptance but cannot yet be described as fully passing due to one administrator consumption-history assertion failure.": (
        "The current automated verification run passed all 181 tests and 1,178 assertions. This evidence supports software-level functional acceptance, but it does not substitute for physical sensor calibration, field-network testing, or usability validation with household users."
    ),
    "In addition, the remaining automated-test failure must be resolved before a final claim of complete regression readiness.": (
        "The completed automated verification supports regression readiness for the tested software behaviors."
    ),
    "and resolution of the remaining automated-test failure.": (
        "and documented physical field validation."
    ),
}

for paragraph in document.paragraphs:
    for run in paragraph.runs:
        for old, new in replacements.items():
            if old in run.text:
                run.text = run.text.replace(old, new)

# Paragraph replacement is necessary where Word split the source sentence into runs.
for paragraph in document.paragraphs:
    text = paragraph.text
    if "137 passing automated tests" in text:
        paragraph.text = (
            "The implemented application provides a complete path from a registered physical meter to role-scoped monitoring and administrative action. "
            "The meter reports measured intervals; the API validates device identity; the database preserves raw and derived values; "
            "services calculate consumption and estimated billing; and the user interfaces expose monitoring, leak, notification, billing, analytics, and reporting functions. "
            "The successful production asset build and 181 passing automated tests with 1,178 assertions demonstrate broad functional coverage."
        )
    elif "cannot yet be described as fully passing" in text:
        paragraph.text = (
            "The current automated verification run passed all 181 tests and 1,178 assertions. This evidence supports software-level functional acceptance, "
            "but it does not substitute for physical sensor calibration, field-network testing, or usability validation with household users."
        )
    elif "remaining automated-test failure" in text:
        paragraph.text = text.replace(
            "In addition, the remaining automated-test failure must be resolved before a final claim of complete regression readiness.",
            "The completed automated verification supports regression readiness for the tested software behaviors.",
        ).replace(
            "and resolution of the remaining automated-test failure.",
            "and documented physical field validation.",
        )

for section in document.sections:
    for paragraph in section.header.paragraphs + section.footer.paragraphs:
        paragraph.text = paragraph.text.replace(
            "SMART WATER MONITORING CONSUMPTION SYSTEM",
            "SMART WATER CONSUMPTION MONITORING SYSTEM",
        )

for paragraph in document.paragraphs:
    paragraph.text = paragraph.text.replace(
        "SMART WATER MONITORING\nCONSUMPTION SYSTEM",
        "SMART WATER CONSUMPTION\nMONITORING SYSTEM",
    )
    if paragraph.text == "Project-based manuscript draft\nCalape, Bohol":
        paragraph.text = "Updated system documentation\nCalape, Bohol"
    elif paragraph.text == "Prepared from the implemented Laravel and ESP32 project\nAugust 2026":
        paragraph.text = "Aligned with the implemented Laravel and ESP32 project\nOctober 2026"

table_replacements = {
    "138 tests; 885 assertions": "181 tests; 1,178 assertions",
    "137 passed; 1 failed": "181 passed; 0 failed",
    "18.017 seconds": "22.390 seconds",
    "The failing case expected one record but received two in an administrator consumption-history response.": "The current test run completed without failed cases.",
    "Current verification status: The project is substantially covered by automated tests, but the suite is not fully green as of 24 August 2026. One consumption-monitoring test must be investigated before claiming complete regression-test acceptance.": "Current verification status: The complete automated suite passed with 181 tests and 1,178 assertions. This verifies the tested software behaviors but does not replace physical sensor calibration or field validation.",
    "Test-suite regression": "Automated regression verification",
    "One current consumption-history test has a record-count mismatch.": "The current automated suite completed without failed cases.",
    "Isolate the query/date fixture, correct the root cause, and rerun the complete suite before acceptance.": "Maintain the automated suite in continuous verification as the system evolves.",
    "Interpretation: The implementation is suitable for continued pilot testing and controlled demonstration. Production acceptance should follow calibration, backup, security, network-recovery, and full regression-test sign-off.": "Interpretation: The implementation is suitable for continued pilot testing and controlled demonstration. Production acceptance should follow calibration, backup, security, network-recovery, official billing review, and documented field validation.",
    "137 of 138 automated tests passed; frontend production build passed.": "181 automated tests with 1,178 assertions passed; frontend production build passed.",
    "Conditional": "Passed for tested software behaviors",
    "Broad coverage is present, but one consumption-history mismatch remains before full acceptance.": "Broad coverage is present. Physical calibration, field-network testing, and usability validation remain required before full deployment acceptance.",
    "Resolve the failing regression test": "Maintain automated regression verification",
    "Investigate why the administrator consumption-history response returns two records when the fixture expects one; confirm date boundaries, filters, and test isolation; then rerun all 138 tests.": "Run the full suite after every material software change and investigate any failed case before release.",
}
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for old, new in table_replacements.items():
                if old in cell.text:
                    cell.text = cell.text.replace(old, new)

# The replacement acceptance text is longer than the source row.  Keep its
# type consistent with the rest of the verification table rather than letting
# Word retain the document's larger default paragraph font.
for table in document.tables:
    for row in table.rows:
        if row.cells and "Regression" in row.cells[0].text:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)

document.save(output)
