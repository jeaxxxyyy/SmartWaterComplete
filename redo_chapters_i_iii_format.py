from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\laragon\capstone\SmartWater_Update\Smart-Water-Consumption-Monitoring-Revised.docx")
OUTPUT = Path(r"C:\laragon\capstone\SmartWater_Update\Smart-Water-Consumption-Monitoring-Updated-Chapters-I-III.docx")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.first_child_found_in("w:tcMar")
    if node is None:
        node = OxmlElement("w:tcMar")
        tc_pr.append(node)
    for side in ("top", "start", "bottom", "end"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), "110")
        edge.set(qn("w:type"), "dxa")
        node.append(edge)


def no_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    return p


def body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.35)
    p.paragraph_format.space_after = Pt(6)
    p.add_run(text)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(3)
    p.add_run("• " + text)
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    for i, (label, width) in enumerate(zip(headers, widths)):
        c = t.rows[0].cells[i]
        c.width = Inches(width)
        shade(c, "1F4E78")
        margins(c)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)
    for row_index, row in enumerate(rows):
        cells = t.add_row().cells
        no_split(t.rows[-1])
        for i, (value, width) in enumerate(zip(row, widths)):
            c = cells[i]
            c.width = Inches(width)
            margins(c)
            if row_index % 2:
                shade(c, "F2F6F9")
            r = c.paragraphs[0].add_run(value)
            r.font.size = Pt(10)
    doc.add_paragraph()
    return t


doc = Document(SOURCE)
body_xml = doc._element.body
for child in list(body_xml):
    if child.tag != qn("w:sectPr"):
        body_xml.remove(child)

for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles["Normal"].font.size = Pt(12)
doc.styles["Normal"].paragraph_format.line_spacing = 1.5
for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True

# This cover identifies the standalone replacement. The body below follows the
# provided Chapter I–III hierarchy exactly, without adding new numbered sections.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(125)
r = p.add_run("SMART WATER CONSUMPTION\nMONITORING SYSTEM")
r.font.name = "Times New Roman"
r.font.size = Pt(18)
r.bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(36)
r = p.add_run("CHAPTERS I TO III")
r.font.size = Pt(14)
r.bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(90)
p.add_run("Updated to follow the approved chapter format\nCalape, Bohol\nOctober 2026").font.size = Pt(12)
doc.add_page_break()

# I. Introduction
heading(doc, "I. Introduction", 1)
heading(doc, "1.1 Project Context", 2)
body(doc, "Households often learn their water use only after an official bill is issued. This delay makes it difficult to connect daily consumption to cost, identify unusual flow early, or decide whether a possible plumbing concern requires attention. Connected water monitoring can provide a shorter feedback cycle by collecting flow information from a household pipeline and presenting it to households and authorized water-service personnel.")
body(doc, "The Smart Water Consumption Monitoring System is a household-scale Internet of Things project for Calape, Bohol. It uses a YF-S201 flow sensor and an M5Stack CoreS3 to observe water flow. The CoreS3 samples the sensor pulse signal, estimates interval flow and volume, shows device status locally, and sends readings through Wi-Fi to a Laravel web application. The application stores accepted readings, derives the monthly consumption total, and provides household and municipal workspaces.")
body(doc, "The updated system includes a municipal device registry and a controlled commissioning process. A device receives a one-time activation code, is bound to a hardware identity and household, and uses a scoped credential for telemetry. Each reading carries a stable identifier. If Wi-Fi or the server is temporarily unavailable, the device retains the pending interval and can resend it safely; the server recognizes a repeated identifier and prevents the same interval from being counted twice.")
heading(doc, "1.2 Purpose and Description", 2)
body(doc, "The purpose of the project is to provide an IoT-based Smart Water Consumption Monitoring System that gives households timely visibility of measured use and gives authorized municipal personnel a role-scoped way to review devices, household records, estimated billing information, notifications, reports, and follow-up activity. It is intended to support awareness and monitoring, not to replace an approved utility meter, an official municipal bill, or a physical inspection.")
body(doc, "The device reads pulses from the YF-S201 on GPIO2 and uses a configurable calibration value to compute interval liters and flow rate. Firmware samples the signal once per second and attempts a telemetry upload approximately every three seconds. The Laravel application validates device identity and household assignment, stores accepted readings, calculates the server-authoritative monthly total, evaluates possible-leak conditions, and applies relevant billing settings. Households can view their own consumption, history, notifications, goals, statements, and service requests. Administrators can manage households, device lifecycle, billing settings, leak records, reports, offline payments, and audit logs.")
heading(doc, "1.3 Objectives", 2)
body(doc, "The general objective is to develop an IoT-based Smart Water Consumption Monitoring System that measures household water use, transmits reliable telemetry to a web application, and provides consumption, notification, and estimated-billing information appropriate to each user role.")
for item in [
    "Measure water-flow pulses through a YF-S201 sensor and convert calibrated pulse counts into interval liters and flow rate using an M5Stack CoreS3.",
    "Provide local device status and upload readings through Wi-Fi using stable identifiers and retained pending intervals for safe retry.",
    "Commission physical meters through a municipal registry, one-time activation, hardware identity binding, and scoped device credentials.",
    "Provide households with live and historical consumption, notifications, goals, statements, billing documents, and service-request functions.",
    "Provide administrators with household, device, billing, leak, report, notification, payment-ledger, service-request, and audit functions.",
    "Identify possible leak conditions from server-configured flow rules and create alerts for investigation without claiming physical diagnosis or automatic water shutoff.",
    "Verify implemented software behavior through automated tests while keeping physical calibration and real-user field validation as separate requirements."
]:
    bullet(doc, item)
heading(doc, "1.4 Scope and Limitations", 2)
body(doc, "The system covers household monitoring through a YF-S201 flow sensor connected to an M5Stack CoreS3. It includes local device display, Wi-Fi telemetry, retry-safe readings, role-scoped household and municipal web workspaces, consumption history, estimated and finalized statement support, configurable rates and subsidies or deductions, notifications, possible-leak records, reports, offline-payment records, service requests, and audit logging.")
body(doc, "The initial firmware setting of 450 pulses per liter is a starting calibration value, not a verified accuracy claim. A measured-volume test is required for each installed sensor and plumbing arrangement. The system depends on power and usable Wi-Fi for timely data upload; an extended outage can delay the dashboard. Possible-leak alerts require inspection and do not close a valve. System estimates and statements must be reconciled with approved municipal rules before being treated as official bills. The project has automated software-test evidence, but sensor accuracy, installation quality, backup restoration, long-term reliability, and unassisted usability still require field validation.")

doc.add_page_break()

# II. Review of Related Literature
heading(doc, "II. Review of Related Literature", 1)
body(doc, "Andrić et al. (2022) describe an IoT approach to smart water use in which connected meters provide more transparent consumption information and support data analysis. The study is relevant to the present project because it treats water metering as a complete information path from physical measurement to understandable consumption data. The present system applies this principle at household scale through Wi-Fi telemetry and a web-based municipal workflow.")
body(doc, "Tasong and Abao (2019) developed a water-monitoring application that combined Hall-effect flow sensing, cloud storage, browser access, and calibration trials. Their work is relevant because it demonstrates that pulse-derived volume must be calibrated rather than assumed. Accordingly, the present project treats its firmware factor as an initial configuration and does not claim field accuracy before a measured-volume calibration test has been completed.")
body(doc, "Bragalli, Neri, and Toth (2019) examined smart-meter-based water-loss assessment in the presence of incomplete readings. This is relevant to the present project because communication problems can cause missing or repeated telemetry. The system responds by preserving unsent intervals locally, sending a stable reading identifier on retry, and applying server-side idempotency so a previously accepted reading is not counted again.")
body(doc, "The project also adopts a controlled device lifecycle. Laravel Sanctum documents token-based API authentication concepts that inform the separation between browser-user access and commissioned device access. The implemented project extends that approach with municipal device registration, one-time activation, hardware binding, scoped credentials, suspension, rotation, and decommissioning. This prevents a generic shared upload key from being the sole identity of a physical meter.")
body(doc, "The reviewed work supports the feasibility of connected water sensing, visual consumption information, calibration-aware measurement, and anomaly-oriented monitoring. The present project integrates those ideas with a local operational workflow: registered device identity, household association, retry-safe telemetry, server-authoritative consumption, configurable billing support, managed possible-leak follow-up, reports, and audit trails. It does not claim a new sensing principle or a fully validated utility replacement.")

doc.add_page_break()

# III. Technical Background
heading(doc, "III. Technical Background", 1)
heading(doc, "3.1 Architecture Design", 2)
heading(doc, "3.1.1 Smart Water Consumption Monitoring System Architecture Design", 3)
body(doc, "The system has four connected layers. The measurement layer is the YF-S201 sensor installed in the household water path. The device layer is the M5Stack CoreS3, which receives the pulse signal through GPIO2, calculates interval values from its calibration setting, retains pending readings, and presents local device information. The application layer is a Laravel service with APIs, a relational database, business services, and audit logging. The user layer consists of household and administrator browser workspaces with records restricted by role and ownership.")
table(doc, ["Layer", "Role in the System", "Implemented Controls"], [
    ("Measurement", "Produces a pulse signal representing water movement.", "YF-S201 flow sensor; measured-volume calibration remains required."),
    ("Device", "Samples pulses, calculates interval values, shows status, and uploads telemetry.", "CoreS3; GPIO2; one-second sampling; three-second upload interval; pending-reading persistence."),
    ("Application", "Validates, stores, derives, notifies, bills, reports, and audits.", "Laravel API; server-side household assignment; stable reading IDs; server total and alert evaluation."),
    ("Household workspace", "Shows only authenticated household records.", "Live consumption, history, statements, notifications, goals, requests, and settings."),
    ("Municipal workspace", "Manages records and operational follow-up.", "Device lifecycle, billing configuration, leak records, reports, offline payments, service requests, and audit logs.")
], [1.25, 2.5, 2.75])
body(doc, "Water passing through the sensor produces pulses. The CoreS3 samples these pulses and uploads a unique interval reading when connectivity is available. The application verifies the commissioned device and its household assignment before storing the reading. It then derives the current-period total from accepted readings, evaluates the applicable billing and possible-leak rules, and makes the resulting state available to the device screen and appropriate web workspace. Household access is limited to the authenticated household; administrators can perform governed configuration and lifecycle actions recorded in audit logs.")
heading(doc, "3.2 Details of Technology Used", 2)
body(doc, "M5Stack CoreS3 ESP32 IoT Development Kit. The CoreS3 is the embedded processing and display device in the project. Its firmware samples the sensor once per second, uses Wi-Fi for telemetry, keeps pending readings for retry, and displays status without exposing credentials. The device does not decide municipal billing thresholds or authoritative monthly totals; these remain server-side decisions.")
body(doc, "YF-S201 Water Flow Sensor. The sensor produces pulse signals as water passes through it. The CoreS3 converts counted pulses into volume and flow information using a configurable factor. The project's initial configuration is 450 pulses per liter. Since calibration depends on the installed sensor and plumbing conditions, measured-volume testing is required before field accuracy can be claimed.")
body(doc, "Laravel Framework and Database Services. Laravel supplies the web application, API controllers, service logic, authentication, database migrations, report generation, and automated tests. WaterConsumptionService validates readings, prevents duplicate interval counting, derives monthly consumption, and creates notifications. BillingStatementService applies effective rates, fixed charges, subsidies or deductions, opening balance, and confirmed offline payments when generating statements.")
body(doc, "Web Application. The browser interface provides household monitoring and municipal administration. Household users can access scoped consumption records, history, statements, notifications, goals, settings, and service requests. Administrators manage household records, physical device lifecycle, billing configuration, leak follow-up, reports, exports, offline-payment receipts, and audit trails. The website does not process online payments; authorized municipal staff record offline payment information.")
body(doc, "Verification Tools. The project uses automated tests to verify device activation, credential checks, duplicate-reading idempotency, telemetry retries, monthly rollover, billing calculations, statements, offline payments, household data isolation, alerts, and selected security controls. The latest verification completed 181 automated tests with 1,178 assertions passing. This evidence supports the tested software behavior only and does not replace calibration, installation, recovery, backup, or real-user validation.")

heading(doc, "References", 1)
for ref in [
    "Andrić, I., Vrsalović, A., Perković, T., Aglić Čuvić, A., & Šolić, P. (2022). IoT approach towards smart water usage. Journal of Cleaner Production, 365, 132792. https://doi.org/10.1016/j.jclepro.2022.132792",
    "Bragalli, C., Neri, M., & Toth, E. (2019). Effectiveness of smart meter-based urban water loss assessment in a real network with synchronous and incomplete readings. Environmental Modelling & Software, 112, 128-142. https://doi.org/10.1016/j.envsoft.2018.10.010",
    "Laravel. (2026). Laravel Sanctum documentation. https://laravel.com/docs/sanctum",
    "M5Stack. (n.d.). CoreS3 documentation. https://docs.m5stack.com/en/core/CoreS3",
    "Tasong, A. C., & Abao, R. P. (2019). Design and development of an IoT application with visual analytics for water consumption monitoring. Procedia Computer Science, 157, 205-213. https://doi.org/10.1016/j.procs.2019.08.159"
]:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    p.paragraph_format.space_after = Pt(5)
    p.add_run(ref)

for section in doc.sections:
    for p in section.header.paragraphs:
        p.text = "SMART WATER CONSUMPTION MONITORING SYSTEM"
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0, 0, 0)

doc.save(OUTPUT)
print(OUTPUT)
